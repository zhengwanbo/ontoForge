import json
import re
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile
from typing import Any, Dict, List, Optional
import xml.etree.ElementTree as ET

from sqlalchemy.orm import Session, selectinload

from app.core.logging import get_logger
from app.models.models import (
    SysDomain,
    SysBusinessType,
    SysOntologyBlueprint,
    SysOntologyEntity,
    SysEntityMapping,
    SysOntologyProperty,
    SysPropertyMapping,
    SysOntologyRelation,
    SysRelationMapping,
    generate_id,
)
from app.services.llm_service import LLMService
from app.services.source_data_service import SourceDataService
from app.services.domain_ontology_generators import build_canonical_model, build_view_plan
from app.services.business_type_service import ensure_default_business_types

logger = get_logger(__name__)


class OntologyGuideService:
    MAX_DOCUMENT_SIZE = 10 * 1024 * 1024
    GUIDE_BLUEPRINT_MAX_TOTAL_COLUMNS = 240
    GUIDE_BLUEPRINT_MAX_TABLES_PER_CHUNK = 4
    GUIDE_LLM_MAX_TOTAL_COLUMNS = 180
    GUIDE_LLM_MIN_COLUMNS_PER_TABLE = 12
    GUIDE_LLM_MAX_COLUMNS_PER_TABLE = 28
    GUIDE_LLM_MAX_SAMPLE_ROWS_PER_TABLE = 1
    GUIDE_LLM_MAX_SAMPLE_COLUMNS_PER_ROW = 12
    RULE_DATA_MAX_RECORDS = 500
    DATABASE_RULE_SAMPLE_LIMIT = 100
    SOURCE_ROLE_CATALOG = {
        "entity_master": "实体主数据",
        "process_history": "过程履历",
        "measurement": "测量结果",
        "rule_catalog": "规则目录",
        "case_library": "案例知识库",
        "event_log": "事件日志",
        "reference_data": "参考字典",
        "other": "其他",
    }
    SEMANTIC_PATTERN_CATALOG = {
        "master-data-linking": {
            "name": "主数据关联",
            "description": "围绕主实体、批次、机种等稳定对象组织基础语义关系。",
        },
        "process-trace": {
            "name": "过程追溯",
            "description": "围绕产品或对象的站位履历、设备、物料、时间构建追溯路径。",
        },
        "measurement-threshold-violation": {
            "name": "测量阈值判定",
            "description": "围绕测量结果与规则目录构建规则判定与异常语义对象。",
        },
        "case-rootcause-action": {
            "name": "案例根因闭环",
            "description": "围绕历史案例、根因与改善动作构建经验复用链路。",
        },
    }

    def __init__(self, db: Session):
        self.db = db
        self.source_service = SourceDataService(db)
        self.llm_service = LLMService(db)
        self._semantic_pattern_catalog = self.SEMANTIC_PATTERN_CATALOG
        self._business_type_context: Dict[str, str] = {}
        self._selected_semantic_type_code: Optional[str] = None

    @staticmethod
    def _normalize_flag(value: Any, default: str = "N") -> str:
        token = str(value or "").strip().upper()
        return token if token in {"Y", "N"} else default

    @staticmethod
    def _normalize_usage_codes(value: Any) -> str:
        raw_items = value if isinstance(value, list) else [value] if value else []
        normalized: List[str] = []
        seen = set()
        for item in raw_items:
            token = str(item or "").strip().lower()
            if token not in {"find", "fetch", "analyze"} or token in seen:
                continue
            seen.add(token)
            normalized.append(token)
        return json.dumps(normalized, ensure_ascii=False)

    @staticmethod
    def _normalize_governance_tags(value: Any) -> str:
        raw_items = value if isinstance(value, list) else [value] if value else []
        normalized: List[str] = []
        seen = set()
        for item in raw_items:
            token = str(item or "").strip()
            if not token or token in seen:
                continue
            seen.add(token)
            normalized.append(token[:50])
        return json.dumps(normalized, ensure_ascii=False)

    @staticmethod
    def _infer_object_type(entity_name: str, entity_desc: str = "") -> Optional[str]:
        sample = f"{entity_name} {entity_desc}".upper()
        fact_tokens = ["ORDER", "EVENT", "RECORD", "LOG", "HISTORY", "TRANSACTION", "REFUND", "MEASURE", "RESULT", "工单", "事件", "记录", "履历", "交易", "订单", "测量", "结果"]
        dim_tokens = ["STORE", "CUSTOMER", "DEVICE", "PRODUCT", "LINE", "MATERIAL", "WAREHOUSE", "REGION", "客户", "门店", "设备", "产品", "产线", "物料", "仓库", "区域"]
        if any(token in sample for token in fact_tokens):
            return "FACT"
        if any(token in sample for token in dim_tokens):
            return "DIM"
        return None

    def _normalize_object_type(self, value: Any, entity_name: str = "", entity_desc: str = "") -> Optional[str]:
        token = str(value or "").strip().upper()
        if token in {"FACT", "DIM"}:
            return token
        return self._infer_object_type(entity_name, entity_desc)

    @staticmethod
    def _extract_property_ref_token(value: Any) -> Optional[str]:
        if isinstance(value, dict):
            property_id = str(value.get("propertyId") or value.get("property_id") or "").strip()
            if property_id:
                return property_id
            entity_name = str(value.get("entityName") or value.get("entity_name") or "").strip()
            property_name = str(value.get("propertyName") or value.get("property_name") or "").strip()
            if entity_name and property_name:
                return f"{entity_name}.{property_name}"
            return None
        token = str(value or "").strip()
        return token or None

    @staticmethod
    def _resolve_property_ref_token(
        token: Optional[str],
        property_by_id: Dict[str, SysOntologyProperty],
        property_by_path: Dict[str, SysOntologyProperty],
    ) -> Optional[SysOntologyProperty]:
        raw = str(token or "").strip()
        if not raw:
            return None
        if raw in property_by_id:
            return property_by_id[raw]
        if "." in raw:
            return property_by_path.get(raw.lower())
        return None

    def _load_domain_semantic_patterns(self, domain: SysDomain) -> None:
        """Load user-maintained semantic patterns for the active business type."""
        ensure_default_business_types(self.db)
        selected_code = (self._selected_semantic_type_code or domain.domain_type or "BUSINESS").strip().upper()
        business_type = self.db.query(SysBusinessType).filter(
            SysBusinessType.type_code == selected_code,
            SysBusinessType.status == "ACTIVE",
        ).first()
        if not business_type:
            if self._selected_semantic_type_code:
                raise ValueError("所选业务语义不存在或已停用")
            self._semantic_pattern_catalog = self.SEMANTIC_PATTERN_CATALOG
            self._business_type_context = {}
            return
        try:
            configured_patterns = json.loads(business_type.semantic_patterns_json or "[]")
        except (TypeError, ValueError):
            configured_patterns = []
        catalog: Dict[str, Dict[str, Any]] = {}
        for item in configured_patterns if isinstance(configured_patterns, list) else []:
            if not isinstance(item, dict):
                continue
            code = (item.get("pattern_code") or item.get("code") or "").strip().lower()
            if not code:
                continue
            catalog[code] = {
                "name": (item.get("pattern_name") or item.get("name") or code).strip(),
                "description": item.get("description") or "",
            }
        self._semantic_pattern_catalog = catalog or self.SEMANTIC_PATTERN_CATALOG
        self._business_type_context = {
            "type_code": business_type.type_code,
            "type_name": business_type.type_name,
            "semantic_desc": business_type.semantic_desc or "",
        }

    def _supports_structured_defect_generation(self, semantic_type_code: Optional[str]) -> bool:
        """Structured pipeline eligibility comes from the managed semantic patterns."""
        ensure_default_business_types(self.db)
        code = (semantic_type_code or "").strip().upper()
        business_type = self.db.query(SysBusinessType).filter(
            SysBusinessType.type_code == code,
            SysBusinessType.status == "ACTIVE",
        ).first()
        if not business_type:
            return False
        try:
            patterns = json.loads(business_type.semantic_patterns_json or "[]")
        except (TypeError, ValueError):
            patterns = []
        pattern_codes = {
            str(item.get("pattern_code") or item.get("code") or "").strip().lower()
            for item in patterns if isinstance(item, dict)
        }
        return {"measurement-threshold-violation", "case-rootcause-action"}.issubset(pattern_codes)

    def parse_uploaded_document(self, file_name: str, content: bytes) -> Dict[str, Any]:
        if not content:
            raise ValueError("上传文件为空")
        if len(content) > self.MAX_DOCUMENT_SIZE:
            raise ValueError("上传文件过大，请控制在 10MB 以内")

        suffix = Path(file_name or "").suffix.lower()
        if suffix in {".txt", ".md", ".csv"}:
            extracted_text = self._decode_text_bytes(content)
            markdown = self._text_to_markdown(extracted_text, preserve_markdown=suffix == ".md")
        elif suffix == ".docx":
            markdown = self._extract_docx_markdown(content)
        elif suffix == ".pdf":
            extracted_text = self._extract_pdf_text(content)
            markdown = self._text_to_markdown(extracted_text)
        else:
            raise ValueError("暂不支持该文件类型，请上传 txt、md、docx 或 pdf 文档")

        normalized_markdown = self._normalize_markdown(markdown)
        if not normalized_markdown:
            raise ValueError("未能从文档中提取到有效文本")

        return {
            "file_name": file_name,
            "file_type": suffix.lstrip("."),
            # text 保留用于兼容既有调用方；Guide 新调用方应优先使用 markdown。
            "text": normalized_markdown,
            "markdown": normalized_markdown,
            "content_format": "markdown",
            "char_count": len(normalized_markdown),
        }

    def parse_uploaded_ddl(self, file_name: str, content: bytes) -> Dict[str, Any]:
        if not content:
            raise ValueError("上传文件为空")
        if len(content) > self.MAX_DOCUMENT_SIZE:
            raise ValueError("上传文件过大，请控制在 10MB 以内")

        suffix = Path(file_name or "").suffix.lower()
        if suffix not in {".sql", ".ddl", ".txt", ".md"}:
            raise ValueError("暂只支持 sql、ddl、txt、md 作为数据库DDL文件")

        ddl_text = self._decode_text_bytes(content)
        normalized_text = self._normalize_document_text(ddl_text)
        if not normalized_text:
            raise ValueError("未能从DDL文件中提取到有效文本")

        tables = self._parse_ddl_tables(normalized_text)
        if not tables:
            raise ValueError("未在DDL文件中识别到 CREATE TABLE 定义")

        return {
            "file_name": file_name,
            "file_type": suffix.lstrip("."),
            "char_count": len(normalized_text),
            "table_count": len(tables),
            "tables": tables,
        }

    def parse_uploaded_rule_data(self, file_name: str, content: bytes) -> Dict[str, Any]:
        if not content:
            raise ValueError("上传文件为空")
        if len(content) > self.MAX_DOCUMENT_SIZE:
            raise ValueError("上传文件过大，请控制在 10MB 以内")

        suffix = Path(file_name or "").suffix.lower()
        if suffix not in {".sql", ".ddl", ".txt", ".md"}:
            raise ValueError("规则数据文件仅支持 sql、ddl、txt、md")

        rule_text = self._decode_text_bytes(content)
        normalized_text = self._normalize_document_text(rule_text)
        if not normalized_text:
            raise ValueError("未能从规则数据文件中提取到有效文本")

        datasets = self._parse_rule_datasets_from_text(normalized_text)
        if not datasets:
            raise ValueError("未识别到可用于缺陷识别的规则数据，当前仅支持解析类似 SPEC_LIMIT 的 INSERT 数据")

        return {
            "file_name": file_name,
            "file_type": suffix.lstrip("."),
            "char_count": len(normalized_text),
            "dataset_count": len(datasets),
            "datasets": datasets,
        }

    async def generate(
        self,
        domain_id: str,
        relation_tables: List[str],
        business_document: str,
        source_id: Optional[str] = None,
        schema: str | None = None,
        table_source_mode: str = "database",
        generation_strategy: Optional[str] = None,
        generation_mode: str = "integrated",
        base_blueprint_id: Optional[str] = None,
        document_generation_phase: str = "ontology",
        design_blueprint_id: Optional[str] = None,
        business_scenario: Optional[str] = None,
        semantic_type_code: Optional[str] = None,
        rule_table_name: Optional[str] = None,
        table_bindings: Optional[List[Dict[str, Any]]] = None,
        ddl_tables: Optional[List[Dict[str, Any]]] = None,
        rule_datasets: Optional[List[Dict[str, Any]]] = None,
        focus_metric_families: Optional[List[str]] = None,
        focus_stations: Optional[List[str]] = None,
        history_case_sources: Optional[List[str]] = None,
        enabled_patterns: Optional[List[str]] = None,
        model_config_id: str | None = None,
        sample_limit: int = 3,
        auto_apply: bool = False,
        overwrite_existing: bool = False,
        created_by: str = "unknown",
    ) -> Dict[str, Any]:
        normalized_generation_mode = (generation_mode or "integrated").strip().lower()
        if normalized_generation_mode == "document_first":
            return await self._generate_document_first(
                domain_id=domain_id,
                business_document=business_document,
                semantic_type_code=semantic_type_code,
                enabled_patterns=enabled_patterns or [],
                model_config_id=model_config_id,
                created_by=created_by,
                document_generation_phase=document_generation_phase,
                design_blueprint_id=design_blueprint_id,
            )
        if normalized_generation_mode not in {"integrated", "data_enrichment"}:
            raise ValueError("不支持的本体生成阶段")
        normalized_generation_strategy = self._normalize_generation_strategy(generation_strategy)
        # 数据补全必须围绕已确认的逻辑基线做语义补全；不能再进入只适用于
        # 结构化制造场景的一次性流水线，否则会丢失基线约束。
        if normalized_generation_mode == "data_enrichment":
            normalized_generation_strategy = "llm_first"
        normalized_business_scenario = self._normalize_business_scenario(business_scenario)
        if normalized_generation_strategy == "structured_domain_pipeline":
            if not self._supports_structured_defect_generation(semantic_type_code):
                raise ValueError("结构化领域生成请从“业务语义管理”中选择已配置“质量规则与缺陷判定”和“根因分析与改善闭环”的业务语义")
            normalized_business_scenario = "DEFECT_ROOTCAUSE"
        self._selected_semantic_type_code = (
            (semantic_type_code or "").strip().upper()
            if normalized_generation_strategy == "llm_first" and (semantic_type_code or "").strip()
            else (semantic_type_code or "").strip().upper()
            if normalized_generation_strategy == "structured_domain_pipeline" and (semantic_type_code or "").strip()
            else None
        )
        guide_context = self._build_guide_strategy_context(
            business_scenario=normalized_business_scenario,
            focus_metric_families=focus_metric_families,
            focus_stations=focus_stations,
            history_case_sources=history_case_sources,
        )

        # LLM 自由生成不使用 SFR/缺陷等固定领域场景模板，而是由当前业务
        # 类型的语义说明及用户选定的语义模式决定本次生成范围。
        if normalized_generation_strategy == "llm_first":
            guide_context["business_scenario"] = "BUSINESS_SEMANTIC"
            guide_context["semantic_type_code"] = self._selected_semantic_type_code

        if normalized_generation_strategy == "structured_domain_pipeline":
            result = await self._generate_with_structured_domain_pipeline(
                domain_id=domain_id,
                relation_tables=relation_tables,
                business_document=business_document,
                source_id=source_id,
                schema=schema,
                table_source_mode=table_source_mode,
                guide_context=guide_context,
                rule_table_name=rule_table_name,
                table_bindings=table_bindings,
                ddl_tables=ddl_tables,
                rule_datasets=rule_datasets,
                enabled_patterns=enabled_patterns,
                model_config_id=model_config_id,
                sample_limit=sample_limit,
                auto_apply=auto_apply,
                overwrite_existing=overwrite_existing,
                created_by=created_by,
            )
            if normalized_generation_mode == "integrated":
                result["generation_mode"] = "integrated"
                self._annotate_integrated_generation_support(result)
                persisted = self._save_blueprint_package(
                    domain_id=domain_id,
                    source_id=result.get("source_id"),
                    schema=result.get("schema"),
                    payload=result,
                    created_by=created_by,
                    status="GENERATED",
                )
                result["blueprint_id"] = persisted["blueprint_id"]
                result["blueprint_version"] = persisted["version_no"]
            return result

        if normalized_generation_mode == "data_enrichment":
            if not base_blueprint_id:
                raise ValueError("数据补全阶段必须选择已确认的逻辑本体基线")
            base_blueprint = self._load_blueprint_payload(base_blueprint_id, domain_id)
            if not base_blueprint:
                raise ValueError("所选逻辑本体基线不存在或不属于当前分析域")
            business_document = f"{business_document}\n\n已确认的逻辑本体基线（必须优先保留并补全映射）：\n{json.dumps({'entities': base_blueprint.get('entities') or [], 'relations': base_blueprint.get('relations') or []}, ensure_ascii=False)}"

        result = await self._generate_with_llm_first(
            domain_id=domain_id,
            relation_tables=relation_tables,
            business_document=business_document,
            source_id=source_id,
            schema=schema,
            table_source_mode=table_source_mode,
            generation_strategy=normalized_generation_strategy,
            guide_context=guide_context,
            rule_table_name=rule_table_name,
            table_bindings=table_bindings,
            ddl_tables=ddl_tables,
            rule_datasets=rule_datasets,
            enabled_patterns=enabled_patterns,
            model_config_id=model_config_id,
            sample_limit=sample_limit,
            # 数据补全和一体化生成都要先标注数据支撑度再持久化。前者需要
            # 合并逻辑基线；后者已同时选择业务文档和数据表，生成对象默认进入
            # 自动映射范围。
            auto_apply=auto_apply if normalized_generation_mode != "data_enrichment" else False,
            overwrite_existing=overwrite_existing,
            created_by=created_by,
            persist_blueprint=normalized_generation_mode not in {"data_enrichment", "integrated"},
        )
        result["generation_mode"] = normalized_generation_mode
        if normalized_generation_mode == "data_enrichment":
            # LLM 的职责是补全物理映射，不能静默删掉已由人工确认的逻辑对象。
            # 先保留本轮识别出的物理字段映射，再用基线补回本轮没有识别到的
            # 逻辑对象；这样同名属性不会被基线中的空来源覆盖。
            merged_baseline = self._merge_blueprint_results([result, base_blueprint])
            result["entities"] = merged_baseline.get("entities") or result.get("entities") or []
            result["relations"] = merged_baseline.get("relations") or result.get("relations") or []
            result["entity_candidates"] = result["entities"]
            result["relation_candidates"] = result["relations"]
            data_support = self._annotate_data_enrichment_support(
                entities=result["entities"],
                relations=result["relations"],
                selected_tables=result.get("selected_tables") or relation_tables,
            )
            result["data_support"] = data_support
            # 自动数据映射只接收已有物理来源依据的对象与边；逻辑候选仍保留
            # 在本体设计中供后续补数或人工映射，不会被自动映射任务扫描。
            supported_entities = [
                item for item in result["entities"]
                if item.get("data_support_status") == "DATA_SUPPORTED"
            ]
            supported_relations = [
                item for item in result["relations"]
                if item.get("data_support_status") == "DATA_SUPPORTED"
            ]
            result["mapping_design"] = self._build_mapping_design(
                entities=supported_entities,
                relations=supported_relations,
                source_role_bindings=result.get("source_role_bindings") or result.get("table_roles") or [],
                semantic_patterns=result.get("semantic_patterns") or [],
            )
            result["base_blueprint_id"] = base_blueprint_id
            result["base_blueprint_version"] = base_blueprint.get("blueprint_version")
            result["baseline_merge_note"] = "已保留已确认逻辑本体；仅标记有实际表和字段依据的对象、属性、边进入自动数据映射范围。"
        elif normalized_generation_mode == "integrated":
            self._annotate_integrated_generation_support(result)

        if normalized_generation_mode in {"data_enrichment", "integrated"}:
            persisted = self._save_blueprint_package(
                domain_id=domain_id,
                source_id=result.get("source_id"),
                schema=result.get("schema"),
                payload=result,
                created_by=created_by,
                status="GENERATED",
            )
            result["blueprint_id"] = persisted["blueprint_id"]
            result["blueprint_version"] = persisted["version_no"]
        return result

    async def _generate_document_first(
        self,
        domain_id: str,
        business_document: str,
        semantic_type_code: Optional[str],
        enabled_patterns: List[str],
        model_config_id: Optional[str],
        created_by: str,
        document_generation_phase: str = "ontology",
        design_blueprint_id: Optional[str] = None,
    ) -> Dict[str, Any]:
        if not (business_document or "").strip():
            raise ValueError("请输入业务说明文档")
        domain = self.db.query(SysDomain).filter(SysDomain.domain_id == domain_id).first()
        if not domain:
            raise ValueError("业务分析域不存在")
        self._selected_semantic_type_code = (semantic_type_code or "").strip().upper() or None
        self._load_domain_semantic_patterns(domain)
        parsed = self._parse_business_document_structured(business_document)
        empty_schema = {"table_count": 0, "tables": []}
        rule_summary = self._build_rule_summary(empty_schema, parsed, [])
        semantic_patterns = self._build_semantic_patterns([], enabled_patterns)
        business_summary = self._build_business_summary(domain, parsed, rule_summary, [])
        phase = (document_generation_phase or "ontology").strip().lower()
        if phase not in {"design", "ontology"}:
            raise ValueError("文档先行生成阶段仅支持 design 或 ontology")
        if phase == "ontology":
            if not design_blueprint_id:
                raise ValueError("请先确认高层本体设计文档，再生成本体对象")
            design_blueprint = self._load_blueprint_payload(design_blueprint_id, domain_id)
            if not design_blueprint or design_blueprint.get("generation_phase") != "design":
                raise ValueError("所选高层本体设计文档不存在或不属于当前分析域")
            design_document = design_blueprint.get("ontology_design_document") or {}
            if not design_document:
                raise ValueError("高层本体设计文档内容为空，请重新生成")
        else:
            design_document = await self.llm_service.generate_ontology_design_document(
                domain=domain, business_summary=business_summary, selected_table_schema=empty_schema,
                rule_summary=rule_summary, table_roles=[], semantic_patterns=semantic_patterns,
                business_document=business_document, config_id=model_config_id, document_first=True,
            )
            result = {
                "domain_id": domain_id, "domain_name": domain.domain_name, "domain_desc": domain.domain_desc,
                "generation_mode": "document_first", "generation_phase": "design", "generation_strategy": "llm_first",
                "lifecycle_stage": "LOGICAL_DESIGN_DRAFT", "generation_scope": "DOCUMENT_FULL", "table_source_mode": "none",
                "business_document": business_document, "business_document_parsed": parsed,
                "business_summary": business_summary, "rule_summary": rule_summary,
                "ontology_design_document": design_document, "semantic_patterns": semantic_patterns,
                "selected_tables": [], "selected_table_schema": empty_schema, "table_roles": [],
                "entities": [], "relations": [], "entity_candidates": [], "relation_candidates": [],
                "mapping_design": {"status": "NOT_STARTED", "message": "请先确认高层本体设计文档；确认后才生成逻辑实体、属性和关系。"},
                "deployment_design": {"status": "NOT_READY", "message": "等待逻辑本体确认和后续数据映射。"},
            }
            persisted = self._save_blueprint_package(
                domain_id=domain_id, source_id=None, schema=None, payload=result,
                created_by=created_by, status="LOGICAL_DESIGN_DRAFT",
            )
            result["blueprint_id"] = persisted["blueprint_id"]
            result["blueprint_version"] = persisted["version_no"]
            return result
        entity_result = await self.llm_service.generate_entity_candidates(
            domain=domain, business_summary=business_summary, ontology_design_document=design_document,
            selected_table_schema=empty_schema, rule_summary=rule_summary, table_roles=[],
            semantic_patterns=semantic_patterns, config_id=model_config_id, document_only=True,
        )
        # 文档先行的目标是完整表达文档业务语义，不能再用 MVP 设计清单
        # 二次裁剪实体；后续可由人工确认删除不需要的候选。
        entities = entity_result.get("entity_candidates") or []
        relation_result = await self.llm_service.generate_relation_candidates(
            domain=domain, business_summary=business_summary, ontology_design_document=design_document,
            rule_summary=rule_summary, table_roles=[], entity_candidates=entities, relation_tables=[],
            semantic_patterns=semantic_patterns, config_id=model_config_id, document_only=True,
        )
        relations = relation_result.get("relation_candidates") or []
        blueprint = self._build_ontology_design(entities, relations)
        result = {
            "domain_id": domain_id, "domain_name": domain.domain_name, "domain_desc": domain.domain_desc,
            "generation_mode": "document_first", "generation_phase": "ontology", "generation_strategy": "llm_first",
            "source_design_blueprint_id": design_blueprint_id,
            "source_design_blueprint_version": design_blueprint.get("blueprint_version"),
            "lifecycle_stage": "LOGICAL_DRAFT", "generation_scope": "DOCUMENT_FULL", "table_source_mode": "none",
            "business_document": business_document, "business_document_parsed": parsed,
            "business_summary": business_summary, "rule_summary": rule_summary,
            "ontology_design_document": design_document, "semantic_patterns": semantic_patterns,
            "selected_tables": [], "selected_table_schema": empty_schema, "table_roles": [],
            "entity_candidates": entities, "relation_candidates": relations,
            "mapping_design": {"status": "NOT_STARTED", "message": "等待后续提供数据表、DDL 或规则数据后再开展映射。"},
            "deployment_design": {"status": "NOT_READY", "message": "逻辑本体确认后，需完成数据映射才可生成部署设计。"},
            **blueprint,
        }
        persisted = self._save_blueprint_package(
            domain_id=domain_id, source_id=None, schema=None, payload=result,
            created_by=created_by, status="LOGICAL_DRAFT",
        )
        result["blueprint_id"] = persisted["blueprint_id"]
        result["blueprint_version"] = persisted["version_no"]
        # SYS_ONTOLOGY_BLUEPRINT.STATUS is VARCHAR2(20) in Oracle.
        # Keep the persisted lifecycle marker within that physical constraint.
        self.mark_blueprint_status(design_blueprint_id, "LOGICAL_DESIGN_DONE")
        return result

    def _load_blueprint_payload(self, blueprint_id: str, domain_id: str) -> Optional[Dict[str, Any]]:
        record = self.db.query(SysOntologyBlueprint).filter(
            SysOntologyBlueprint.blueprint_id == blueprint_id,
            SysOntologyBlueprint.domain_id == domain_id,
        ).first()
        if not record:
            return None
        try:
            payload = json.loads(record.blueprint_json or "{}")
        except Exception:
            return None
        if not isinstance(payload, dict):
            return None
        return {**payload, "blueprint_id": record.blueprint_id, "blueprint_version": record.version_no, "blueprint_status": record.status}

    async def _generate_with_structured_domain_pipeline(
        self,
        domain_id: str,
        relation_tables: List[str],
        business_document: str,
        source_id: Optional[str],
        schema: Optional[str],
        table_source_mode: str,
        guide_context: Dict[str, Any],
        rule_table_name: Optional[str],
        table_bindings: Optional[List[Dict[str, Any]]],
        ddl_tables: Optional[List[Dict[str, Any]]],
        rule_datasets: Optional[List[Dict[str, Any]]],
        enabled_patterns: Optional[List[str]],
        model_config_id: Optional[str],
        sample_limit: int,
        auto_apply: bool,
        overwrite_existing: bool,
        created_by: str,
    ) -> Dict[str, Any]:
        prepared = self._prepare_guide_generation_context(
            domain_id=domain_id,
            relation_tables=relation_tables,
            business_document=business_document,
            source_id=source_id,
            schema=schema,
            table_source_mode=table_source_mode,
            guide_context=guide_context,
            rule_table_name=rule_table_name,
            table_bindings=table_bindings,
            ddl_tables=ddl_tables,
            rule_datasets=rule_datasets,
            enabled_patterns=enabled_patterns,
            model_config_id=model_config_id,
            sample_limit=sample_limit,
        )
        business_document_parsed = prepared.get("business_document_parsed") or self._parse_business_document_structured(business_document)
        rule_analysis = self._analyze_spec_limit_rule_data(
            rule_datasets=prepared.get("normalized_rule_datasets") or [],
            rule_summary=prepared.get("rule_summary") or {},
            guide_context=guide_context,
        )
        schema_analysis = self._analyze_source_schema_keywords(
            selected_table_schema=prepared.get("selected_table_schema") or {},
            table_roles=prepared.get("table_roles") or [],
            source_role_bindings=prepared.get("source_role_bindings") or [],
            business_document_parsed=business_document_parsed,
        )
        document_facts = self._extract_document_facts(
            business_document_parsed=business_document_parsed,
            guide_context=guide_context,
            rule_analysis=rule_analysis,
            schema_analysis=schema_analysis,
        )
        focus_scope = self._build_focus_scope(
            guide_context=guide_context,
            document_facts=document_facts,
            rule_analysis=rule_analysis,
            schema_analysis=schema_analysis,
        )
        metric_semantics = self._derive_metric_semantics(
            rule_analysis=rule_analysis,
            business_document_parsed=business_document_parsed,
        )

        business_scenario = guide_context.get("business_scenario") or document_facts.get("business_scenario")
        guide_context = {
            **guide_context,
            "business_scenario": business_scenario,
        }
        analysis_context = {
            "guide_context": guide_context,
            "document_facts": document_facts,
            "rule_analysis": rule_analysis,
            "schema_analysis": schema_analysis,
            "focus_scope": focus_scope,
            "metric_semantics": metric_semantics,
            "selected_table_schema": prepared.get("selected_table_schema") or {},
            "table_roles": prepared.get("table_roles") or [],
            "source_role_bindings": prepared.get("source_role_bindings") or [],
        }
        canonical_model = build_canonical_model(analysis_context)
        view_plan = build_view_plan(analysis_context, canonical_model)
        ontology_design_document = await self.llm_service.generate_structured_ontology_scope_document(
            domain=prepared["domain"],
            business_summary=prepared.get("business_summary") or {},
            document_facts=document_facts,
            rule_analysis=rule_analysis,
            schema_analysis=schema_analysis,
            focus_scope=focus_scope,
            canonical_model=canonical_model,
            config_id=model_config_id,
        )
        canonical_enrichment = await self.llm_service.enrich_structured_canonical_model(
            domain=prepared["domain"],
            business_summary=prepared.get("business_summary") or {},
            rule_analysis=rule_analysis,
            schema_analysis=schema_analysis,
            canonical_model=canonical_model,
            selected_table_schema=prepared.get("selected_table_schema") or {},
            table_roles=prepared.get("table_roles") or [],
            config_id=model_config_id,
        )
        enriched_canonical_model = self._merge_canonical_model_enrichment(
            canonical_model=canonical_model,
            enrichment=canonical_enrichment,
        )
        view_plan_enrichment = await self.llm_service.enrich_structured_view_plan(
            domain=prepared["domain"],
            business_document=business_document,
            canonical_model=enriched_canonical_model,
            view_plan=view_plan,
            selected_table_schema=prepared.get("llm_selected_table_schema") or prepared.get("selected_table_schema") or {},
            config_id=model_config_id,
        )
        enriched_view_plan = self._merge_view_plan_enrichment(
            view_plan=view_plan,
            enrichment=view_plan_enrichment,
        )
        mapping_design = self._build_mapping_design(
            entities=enriched_canonical_model.get("entities") or [],
            relations=enriched_canonical_model.get("relations") or [],
            source_role_bindings=prepared.get("source_role_bindings") or [],
            semantic_patterns=prepared.get("semantic_patterns") or [],
        )
        deployment_design = self._build_deployment_design_from_view_plan(
            canonical_model=enriched_canonical_model,
            view_plan=enriched_view_plan,
        )

        result: Dict[str, Any] = {
            "domain_id": prepared["domain"].domain_id,
            "domain_name": prepared["domain"].domain_name,
            "domain_desc": prepared["domain"].domain_desc,
            "source_id": source_id,
            "schema": schema,
            "table_source_mode": prepared.get("normalized_source_mode"),
            "business_scenario": business_scenario,
            "generation_strategy": "structured_domain_pipeline",
            "guide_context": guide_context,
            "document_facts": document_facts,
            "rule_analysis": rule_analysis,
            "schema_analysis": schema_analysis,
            "focus_scope": focus_scope,
            "metric_semantics": metric_semantics,
            "selected_tables": prepared.get("normalized_tables") or [],
            "selected_rule_table": prepared.get("normalized_rule_table_name") or None,
            "business_document": business_document,
            "business_document_parsed": business_document_parsed,
            "selected_table_schema": prepared.get("selected_table_schema") or {},
            "rule_summary": prepared.get("rule_summary") or {},
            "spec_limit_summary": prepared.get("spec_limit_summary") or {},
            "rule_datasets": prepared.get("normalized_rule_datasets") or [],
            "business_summary": prepared.get("business_summary") or {},
            "ontology_design_document": ontology_design_document,
            "table_roles": prepared.get("table_roles") or [],
            "entity_candidates": enriched_canonical_model.get("entities") or [],
            "relation_candidates": enriched_canonical_model.get("relations") or [],
            "source_role_bindings": prepared.get("source_role_bindings") or [],
            "semantic_patterns": prepared.get("semantic_patterns") or [],
            "ontology_generation_context": {
                "table_count": len(prepared.get("blueprint_table_details") or []),
                "chunk_count": 0,
                "total_columns": sum(item.get("total_columns") or 0 for item in (prepared.get("blueprint_table_details") or [])),
                "mode": "structured_domain_pipeline",
                "input_budget_tokens": (prepared.get("runtime_limits") or {}).get("input_budget_tokens"),
                "context_window_tokens": (prepared.get("runtime_limits") or {}).get("context_window_tokens"),
                "raw_entity_candidate_count": len(enriched_canonical_model.get("entities") or []),
                "filtered_entity_candidate_count": len(enriched_canonical_model.get("entities") or []),
                "raw_relation_candidate_count": len(enriched_canonical_model.get("relations") or []),
                "filtered_relation_candidate_count": len(enriched_canonical_model.get("relations") or []),
            },
            "llm_context_summary": prepared.get("llm_context_summary") or {},
            "llm_selected_table_schema": prepared.get("llm_selected_table_schema") or {},
            "canonical_model": enriched_canonical_model,
            "view_plan": enriched_view_plan,
            "entities": enriched_canonical_model.get("entities") or [],
            "relations": enriched_canonical_model.get("relations") or [],
            "mapping_design": mapping_design,
            "deployment_design": deployment_design,
        }
        llm_enrichment = {
            "structured_analysis_ready": True,
            "structured_analysis_stage": "phase4",
            "canonical_model_generation_mode": "domain_template",
            "view_plan_generation_mode": "domain_template",
            "scope_document_generation_mode": ontology_design_document.get("generation_mode"),
            "canonical_enrichment_generation_mode": canonical_enrichment.get("generation_mode"),
            "view_plan_enrichment_generation_mode": view_plan_enrichment.get("generation_mode"),
            "scope_document_model": ontology_design_document.get("model"),
            "canonical_enrichment_model": canonical_enrichment.get("model"),
            "view_plan_enrichment_model": view_plan_enrichment.get("model"),
        }
        llm_enrichment["structured_analysis_ready"] = True
        result["llm_enrichment"] = llm_enrichment
        if auto_apply:
            blueprint = {
                "entities": result.get("entities") or [],
                "relations": result.get("relations") or [],
            }
            result["apply_result"] = self.apply_blueprint(
                domain_id=domain_id,
                blueprint=blueprint,
                overwrite_existing=overwrite_existing,
                created_by=created_by,
            )
        persisted = self._save_blueprint_package(
            domain_id=domain_id,
            source_id=source_id,
            schema=schema,
            payload=result,
            created_by=created_by,
            status="APPLIED" if auto_apply else "GENERATED",
        )
        result["blueprint_id"] = persisted["blueprint_id"]
        result["blueprint_version"] = persisted["version_no"]
        return result

    async def _generate_with_llm_first(
        self,
        domain_id: str,
        relation_tables: List[str],
        business_document: str,
        source_id: Optional[str] = None,
        schema: str | None = None,
        table_source_mode: str = "database",
        generation_strategy: str = "llm_first",
        guide_context: Optional[Dict[str, Any]] = None,
        rule_table_name: Optional[str] = None,
        table_bindings: Optional[List[Dict[str, Any]]] = None,
        ddl_tables: Optional[List[Dict[str, Any]]] = None,
        rule_datasets: Optional[List[Dict[str, Any]]] = None,
        enabled_patterns: Optional[List[str]] = None,
        model_config_id: str | None = None,
        sample_limit: int = 3,
        auto_apply: bool = False,
        overwrite_existing: bool = False,
        created_by: str = "unknown",
        persist_blueprint: bool = True,
    ) -> Dict[str, Any]:
        guide_context = guide_context or self._build_guide_strategy_context(
            business_scenario=None,
            focus_metric_families=None,
            focus_stations=None,
            history_case_sources=None,
        )
        domain = self.db.query(SysDomain).filter(SysDomain.domain_id == domain_id).first()
        if not domain:
            raise ValueError("业务分析域不存在")
        self._load_domain_semantic_patterns(domain)

        normalized_bindings = self._normalize_table_bindings(relation_tables, table_bindings or [])
        normalized_tables = [item["table_name"] for item in normalized_bindings]
        normalized_rule_table_name = (rule_table_name or "").strip().upper()
        normalized_source_mode = (table_source_mode or "database").strip().lower()
        ddl_table_details = self._normalize_ddl_tables(ddl_tables or [])
        ddl_table_index = {
            (item.get("table_name") or "").upper(): item
            for item in ddl_table_details
            if item.get("table_name")
        }

        if not normalized_tables:
            raise ValueError("请至少选择一张业务关系表")
        if not (business_document or "").strip():
            raise ValueError("请输入业务说明文档")
        if normalized_source_mode not in {"database", "ddl"}:
            raise ValueError("不支持的表信息来源类型")
        if normalized_source_mode == "database" and not source_id:
            raise ValueError("数据库表模式下必须选择数据库连接")
        if normalized_source_mode == "ddl" and not ddl_table_index:
            raise ValueError("DDL 文件模式下请先上传并解析数据库DDL文件")

        prioritized_bindings = self._prioritize_table_bindings_for_guide(
            bindings=normalized_bindings,
            business_document=business_document,
        )
        logger.info(
            "Guide selected tables prioritized: selected=%s prioritized_for_detail=%s",
            normalized_tables,
            [item.get("table_name") for item in prioritized_bindings],
        )
        table_details = []
        for binding in prioritized_bindings:
            table_name = binding["table_name"]
            if normalized_source_mode == "ddl":
                detail = ddl_table_index.get(table_name)
                if not detail:
                    raise ValueError(f"DDL 文件中未找到已选表：{table_name}")
                detail = {**detail}
            else:
                detail = self.source_service.get_remote_table_detail(
                    source_id=source_id,
                    table_name=table_name,
                    schema=schema,
                    sample_limit=max(1, min(sample_limit, 5)),
                )
            detail["source_role"] = self._resolve_binding_role(table_name, normalized_bindings)
            table_details.append(detail)

        rule_table_detail = None
        if normalized_source_mode == "database" and source_id and normalized_rule_table_name:
            if normalized_rule_table_name in {item.get("table_name") for item in table_details}:
                rule_table_detail = next(
                    (item for item in table_details if item.get("table_name") == normalized_rule_table_name),
                    None,
                )
            else:
                rule_table_detail = self.source_service.get_remote_table_detail(
                    source_id=source_id,
                    table_name=normalized_rule_table_name,
                    schema=schema,
                    sample_limit=max(self.DATABASE_RULE_SAMPLE_LIMIT, max(1, min(sample_limit, 5))),
                )
                rule_table_detail["source_role"] = "rule_catalog"
                table_details.append(rule_table_detail)

        business_document_parsed = self._parse_business_document_structured(business_document)
        selected_table_schema = self._build_selected_table_schema(table_details)
        normalized_rule_datasets = self._normalize_rule_datasets(rule_datasets or [])
        if normalized_source_mode == "database" and source_id:
            normalized_rule_datasets.extend(
                self._load_database_rule_datasets(
                    source_id=source_id,
                    schema=schema,
                    relation_tables=normalized_tables,
                    explicit_rule_table_name=normalized_rule_table_name,
                    table_details=table_details,
                )
            )
            normalized_rule_datasets = self._normalize_rule_datasets(normalized_rule_datasets)
        rule_summary = self._build_rule_summary(
            selected_table_schema,
            business_document_parsed,
            normalized_rule_datasets,
        )
        spec_limit_summary = rule_summary.get("spec_limit_summary") or {}
        table_roles = self._build_table_roles(table_details, rule_summary)
        source_role_bindings = self._build_source_role_bindings(table_details)
        blueprint_table_details = self._build_blueprint_table_details(table_details)
        runtime_limits = self.llm_service.get_runtime_limits(config_id=model_config_id)
        blueprint_chunks = self._chunk_blueprint_tables(
            domain=domain,
            business_document=business_document,
            table_details=blueprint_table_details,
            source_role_bindings=source_role_bindings,
            enabled_patterns=enabled_patterns or [],
            input_budget_tokens=runtime_limits.get("input_budget_tokens"),
        )
        llm_table_details, llm_context_summary = self._build_llm_table_details(table_details)
        llm_selected_table_schema = self._build_selected_table_schema(
            llm_table_details,
            for_llm=True,
        )
        logger.info(
            "Guide selected table schema compacted for design document: table_count=%s full_columns=%s compact_columns=%s",
            llm_selected_table_schema.get("table_count"),
            sum(int(item.get("column_count") or 0) for item in (selected_table_schema.get("tables") or [])),
            sum(int(item.get("column_count") or 0) for item in (llm_selected_table_schema.get("tables") or [])),
        )
        llm_context_summary["runtime_limits"] = runtime_limits
        semantic_patterns = self._build_semantic_patterns(
            source_role_bindings=source_role_bindings,
            enabled_patterns=enabled_patterns or [],
        )
        business_summary = self._build_business_summary(
            domain=domain,
            business_document_parsed=business_document_parsed,
            rule_summary=rule_summary,
            table_roles=table_roles,
        )
        ontology_design_document = await self.llm_service.generate_ontology_design_document(
            domain=domain,
            business_summary=business_summary,
            selected_table_schema=llm_selected_table_schema,
            rule_summary=rule_summary,
            table_roles=table_roles,
            business_document=business_document,
            semantic_patterns=semantic_patterns,
            config_id=model_config_id,
        )

        entity_candidate_results: List[Dict[str, Any]] = []
        for chunk in blueprint_chunks:
            chunk_bindings = self._filter_source_role_bindings(source_role_bindings, chunk)
            chunk_patterns = self._build_semantic_patterns(
                source_role_bindings=chunk_bindings,
                enabled_patterns=enabled_patterns or [],
            )
            chunk_results = await self._generate_entity_candidate_chunk_results(
                domain=domain,
                business_summary=business_summary,
                ontology_design_document=ontology_design_document,
                rule_summary=rule_summary,
                selected_table_schema={"table_count": len(chunk), "tables": chunk},
                chunk=chunk,
                enabled_patterns=enabled_patterns or [],
                config_id=model_config_id,
                table_roles=[item for item in table_roles if (item.get("table_name") or "").upper() in {(x.get("table_name") or "").upper() for x in chunk}],
                semantic_patterns=chunk_patterns,
            )
            entity_candidate_results.extend(chunk_results)
        # 仓库、门店等主数据是供应链分析的稳定锚点。即使 LLM 将首期
        # 范围收敛到交易或过程对象，也不能把这类已选且有稳定主键的表静默丢弃。
        mandatory_master_entities = self._build_mandatory_master_entity_candidates(
            selected_table_schema=selected_table_schema,
            table_roles=table_roles,
        )
        raw_entity_candidates = self._merge_entity_candidate_results(
            entity_candidate_results + ([{"entity_candidates": mandatory_master_entities}] if mandatory_master_entities else [])
        )
        entity_candidates = self._filter_entity_candidates_by_design_document(
            raw_entity_candidates,
            ontology_design_document,
        )
        relation_candidates_result = await self.llm_service.generate_relation_candidates(
            domain=domain,
            business_summary=business_summary,
            ontology_design_document=ontology_design_document,
            rule_summary=rule_summary,
            table_roles=table_roles,
            entity_candidates=entity_candidates,
            relation_tables=llm_table_details,
            semantic_patterns=semantic_patterns,
            config_id=model_config_id,
        )
        raw_relation_candidates = relation_candidates_result.get("relation_candidates") or []
        relation_candidates = self._filter_relation_candidates_by_design_document(
            raw_relation_candidates,
            entity_candidates,
            ontology_design_document,
        )
        blueprint = self._build_ontology_design(entity_candidates, relation_candidates)
        mapping_design = self._build_mapping_design(
            entities=blueprint.get("entities") or [],
            relations=blueprint.get("relations") or [],
            source_role_bindings=source_role_bindings,
            semantic_patterns=semantic_patterns,
        )
        base_deployment_design = self._build_deployment_design(
            domain=domain,
            entities=blueprint.get("entities") or [],
            relations=blueprint.get("relations") or [],
            source_role_bindings=source_role_bindings,
            semantic_patterns=semantic_patterns,
        )
        deployment_design = await self.llm_service.generate_semantic_deployment_design(
            domain=domain,
            business_document=business_document,
            entities=blueprint.get("entities") or [],
            relations=blueprint.get("relations") or [],
            relation_tables=llm_table_details,
            source_role_bindings=source_role_bindings,
            semantic_patterns=semantic_patterns,
            base_deployment_design=base_deployment_design,
            config_id=model_config_id,
        )

        result: Dict[str, Any] = {
            "domain_id": domain.domain_id,
            "domain_name": domain.domain_name,
            "domain_desc": domain.domain_desc,
            "source_id": source_id,
            "schema": schema,
            "table_source_mode": normalized_source_mode,
            "generation_strategy": generation_strategy,
            "business_scenario": guide_context.get("business_scenario"),
            "guide_context": guide_context,
            "document_facts": {},
            "rule_analysis": {},
            "schema_analysis": {},
            "focus_scope": {
                "focus_metric_families": guide_context.get("focus_metric_families") or [],
                "focus_stations": guide_context.get("focus_stations") or [],
                "history_case_sources": guide_context.get("history_case_sources") or [],
            },
            "metric_semantics": {},
            "canonical_model": {},
            "view_plan": {},
            "llm_enrichment": {
                "entity_candidate_generation_mode": "llm_first",
                "relation_candidate_generation_mode": "llm_first",
                "deployment_generation_mode": "llm_first",
            },
            "selected_tables": normalized_tables,
            "selected_rule_table": normalized_rule_table_name or None,
            "business_document": business_document,
            "business_document_parsed": business_document_parsed,
            "selected_table_schema": selected_table_schema,
            "rule_summary": rule_summary,
            "spec_limit_summary": spec_limit_summary,
            "rule_datasets": normalized_rule_datasets,
            "business_summary": business_summary,
            "ontology_design_document": ontology_design_document,
            "table_roles": table_roles,
            "entity_candidates": entity_candidates,
            "relation_candidates": relation_candidates,
            "source_role_bindings": source_role_bindings,
            "semantic_patterns": semantic_patterns,
            "ontology_generation_context": {
                "table_count": len(blueprint_table_details),
                "chunk_count": len(blueprint_chunks),
                "total_columns": sum(item.get("total_columns") or 0 for item in blueprint_table_details),
                "mode": "full_columns_chunked",
                "input_budget_tokens": runtime_limits.get("input_budget_tokens"),
                "context_window_tokens": runtime_limits.get("context_window_tokens"),
                "raw_entity_candidate_count": len(raw_entity_candidates),
                "filtered_entity_candidate_count": len(entity_candidates),
                "mandatory_master_entity_count": len(mandatory_master_entities),
                "raw_relation_candidate_count": len(raw_relation_candidates),
                "filtered_relation_candidate_count": len(relation_candidates),
            },
            "llm_context_summary": llm_context_summary,
            "llm_selected_table_schema": llm_selected_table_schema,
            "mapping_design": mapping_design,
            "deployment_design": deployment_design,
            **blueprint,
        }

        if auto_apply:
            result["apply_result"] = self.apply_blueprint(
                domain_id=domain.domain_id,
                blueprint=blueprint,
                overwrite_existing=overwrite_existing,
                created_by=created_by,
            )

        if persist_blueprint:
            persisted = self._save_blueprint_package(
                domain_id=domain.domain_id,
                source_id=source_id,
                schema=schema,
                payload=result,
                created_by=created_by,
                status="APPLIED" if auto_apply else "GENERATED",
            )
            result["blueprint_id"] = persisted["blueprint_id"]
            result["blueprint_version"] = persisted["version_no"]

        return result

    def _prepare_guide_generation_context(
        self,
        domain_id: str,
        relation_tables: List[str],
        business_document: str,
        source_id: Optional[str],
        schema: Optional[str],
        table_source_mode: str,
        guide_context: Dict[str, Any],
        rule_table_name: Optional[str],
        table_bindings: Optional[List[Dict[str, Any]]],
        ddl_tables: Optional[List[Dict[str, Any]]],
        rule_datasets: Optional[List[Dict[str, Any]]],
        enabled_patterns: Optional[List[str]],
        model_config_id: Optional[str],
        sample_limit: int,
    ) -> Dict[str, Any]:
        domain = self.db.query(SysDomain).filter(SysDomain.domain_id == domain_id).first()
        if not domain:
            raise ValueError("业务分析域不存在")
        self._load_domain_semantic_patterns(domain)

        normalized_bindings = self._normalize_table_bindings(relation_tables, table_bindings or [])
        normalized_tables = [item["table_name"] for item in normalized_bindings]
        normalized_rule_table_name = (rule_table_name or "").strip().upper()
        normalized_source_mode = (table_source_mode or "database").strip().lower()
        ddl_table_details = self._normalize_ddl_tables(ddl_tables or [])
        ddl_table_index = {
            (item.get("table_name") or "").upper(): item
            for item in ddl_table_details
            if item.get("table_name")
        }
        if not normalized_tables:
            raise ValueError("请至少选择一张业务关系表")
        if not (business_document or "").strip():
            raise ValueError("请输入业务说明文档")
        if normalized_source_mode not in {"database", "ddl"}:
            raise ValueError("不支持的表信息来源类型")
        if normalized_source_mode == "database" and not source_id:
            raise ValueError("数据库表模式下必须选择数据库连接")
        if normalized_source_mode == "ddl" and not ddl_table_index:
            raise ValueError("DDL 文件模式下请先上传并解析数据库DDL文件")

        prioritized_bindings = self._prioritize_table_bindings_for_guide(
            bindings=normalized_bindings,
            business_document=business_document,
        )
        table_details = []
        for binding in prioritized_bindings:
            table_name = binding["table_name"]
            if normalized_source_mode == "ddl":
                detail = ddl_table_index.get(table_name)
                if not detail:
                    raise ValueError(f"DDL 文件中未找到已选表：{table_name}")
                detail = {**detail}
            else:
                detail = self.source_service.get_remote_table_detail(
                    source_id=source_id,
                    table_name=table_name,
                    schema=schema,
                    sample_limit=max(1, min(sample_limit, 5)),
                )
            detail["source_role"] = self._resolve_binding_role(table_name, normalized_bindings)
            table_details.append(detail)

        if normalized_source_mode == "database" and source_id and normalized_rule_table_name:
            if normalized_rule_table_name not in {item.get("table_name") for item in table_details}:
                rule_table_detail = self.source_service.get_remote_table_detail(
                    source_id=source_id,
                    table_name=normalized_rule_table_name,
                    schema=schema,
                    sample_limit=max(self.DATABASE_RULE_SAMPLE_LIMIT, max(1, min(sample_limit, 5))),
                )
                rule_table_detail["source_role"] = "rule_catalog"
                table_details.append(rule_table_detail)

        business_document_parsed = self._parse_business_document_structured(business_document)
        selected_table_schema = self._build_selected_table_schema(table_details)
        normalized_rule_datasets = self._normalize_rule_datasets(rule_datasets or [])
        if normalized_source_mode == "database" and source_id:
            normalized_rule_datasets.extend(
                self._load_database_rule_datasets(
                    source_id=source_id,
                    schema=schema,
                    relation_tables=normalized_tables,
                    explicit_rule_table_name=normalized_rule_table_name,
                    table_details=table_details,
                )
            )
            normalized_rule_datasets = self._normalize_rule_datasets(normalized_rule_datasets)
        rule_summary = self._build_rule_summary(
            selected_table_schema,
            business_document_parsed,
            normalized_rule_datasets,
        )
        spec_limit_summary = rule_summary.get("spec_limit_summary") or {}
        table_roles = self._build_table_roles(table_details, rule_summary)
        source_role_bindings = self._build_source_role_bindings(table_details)
        blueprint_table_details = self._build_blueprint_table_details(table_details)
        runtime_limits = self.llm_service.get_runtime_limits(config_id=model_config_id)
        llm_table_details, llm_context_summary = self._build_llm_table_details(table_details)
        llm_selected_table_schema = self._build_selected_table_schema(
            llm_table_details,
            for_llm=True,
        )
        llm_context_summary["runtime_limits"] = runtime_limits
        semantic_patterns = self._build_semantic_patterns(
            source_role_bindings=source_role_bindings,
            enabled_patterns=enabled_patterns or [],
        )
        business_summary = self._build_business_summary(
            domain=domain,
            business_document_parsed=business_document_parsed,
            rule_summary=rule_summary,
            table_roles=table_roles,
        )
        return {
            "domain": domain,
            "guide_context": guide_context,
            "normalized_bindings": normalized_bindings,
            "normalized_tables": normalized_tables,
            "normalized_rule_table_name": normalized_rule_table_name,
            "normalized_source_mode": normalized_source_mode,
            "table_details": table_details,
            "business_document_parsed": business_document_parsed,
            "selected_table_schema": selected_table_schema,
            "normalized_rule_datasets": normalized_rule_datasets,
            "rule_summary": rule_summary,
            "spec_limit_summary": spec_limit_summary,
            "table_roles": table_roles,
            "source_role_bindings": source_role_bindings,
            "blueprint_table_details": blueprint_table_details,
            "runtime_limits": runtime_limits,
            "llm_table_details": llm_table_details,
            "llm_context_summary": llm_context_summary,
            "llm_selected_table_schema": llm_selected_table_schema,
            "semantic_patterns": semantic_patterns,
            "business_summary": business_summary,
        }

    def _normalize_generation_strategy(self, generation_strategy: Optional[str]) -> str:
        normalized = str(generation_strategy or "").strip().lower()
        if normalized in {"structured_domain_pipeline", "llm_first"}:
            return normalized
        return "structured_domain_pipeline"

    @staticmethod
    def _normalize_business_scenario(business_scenario: Optional[str]) -> Optional[str]:
        """Unify historical SFR/defect choices into one configurable semantic."""
        normalized = str(business_scenario or "").strip().upper()
        if normalized in {"SFR_ROOTCAUSE", "DEFECT_ANALYSIS", "DEFECT_ROOTCAUSE"}:
            return "DEFECT_ROOTCAUSE"
        return normalized or None

    def _build_guide_strategy_context(
        self,
        business_scenario: Optional[str],
        focus_metric_families: Optional[List[str]],
        focus_stations: Optional[List[str]],
        history_case_sources: Optional[List[str]],
    ) -> Dict[str, Any]:
        return {
            "business_scenario": (business_scenario or "").strip() or None,
            "focus_metric_families": self._normalize_upper_string_list(focus_metric_families),
            "focus_stations": self._normalize_upper_string_list(focus_stations),
            "history_case_sources": self._normalize_upper_string_list(history_case_sources),
        }

    def _normalize_upper_string_list(self, values: Optional[List[str]]) -> List[str]:
        normalized: List[str] = []
        seen = set()
        for value in values or []:
            item = str(value or "").strip().upper()
            if not item or item in seen:
                continue
            seen.add(item)
            normalized.append(item)
        return normalized

    def _extract_distinct_regex_tokens(self, text: str, pattern: str) -> List[str]:
        matches = re.findall(pattern, text or "", re.IGNORECASE)
        normalized: List[str] = []
        seen = set()
        for match in matches:
            item = str(match or "").strip().upper()
            if not item or item in seen:
                continue
            seen.add(item)
            normalized.append(item)
        return normalized

    def _count_station_prefix_columns(self, column_names: List[str]) -> int:
        prefixes = set()
        suffixes = [
            "_MC_ID", "_OP_ID", "_LOT_ID", "_SUB_LOT_ID", "_INPUT_TIME",
            "_OUTPUT_TIME", "_START_TIME", "_END_TIME", "_TOOLING",
        ]
        for column_name in column_names:
            upper_name = str(column_name or "").upper()
            for suffix in suffixes:
                if upper_name.endswith(suffix):
                    prefixes.add(upper_name[: -len(suffix)])
                    break
        return len([item for item in prefixes if item])

    async def _generate_entity_candidate_chunk_results(
        self,
        domain: SysDomain,
        business_summary: Dict[str, Any],
        ontology_design_document: Dict[str, Any],
        rule_summary: Dict[str, Any],
        selected_table_schema: Dict[str, Any],
        chunk: List[Dict[str, Any]],
        enabled_patterns: List[str],
        config_id: Optional[str],
        table_roles: List[Dict[str, Any]],
        semantic_patterns: List[Dict[str, Any]],
    ) -> List[Dict[str, Any]]:
        try:
            result = await self.llm_service.generate_entity_candidates(
                domain=domain,
                business_summary=business_summary,
                ontology_design_document=ontology_design_document,
                selected_table_schema=selected_table_schema,
                rule_summary=rule_summary,
                table_roles=table_roles,
                semantic_patterns=semantic_patterns,
                config_id=config_id,
            )
            return [result]
        except ValueError as exc:
            if "上下文超限风险" not in str(exc):
                raise
            logger.warning(
                "Guide chunk exceeds LLM budget, fallback split: tables=%s reason=%s",
                [item.get("table_name") for item in chunk],
                str(exc),
            )

            if len(chunk) > 1:
                mid = max(1, len(chunk) // 2)
                left = chunk[:mid]
                right = chunk[mid:]
                left_roles = [item for item in table_roles if (item.get("table_name") or "").upper() in {(x.get("table_name") or "").upper() for x in left}]
                right_roles = [item for item in table_roles if (item.get("table_name") or "").upper() in {(x.get("table_name") or "").upper() for x in right}]
                left_patterns = self._build_semantic_patterns(
                    source_role_bindings=left_roles,
                    enabled_patterns=enabled_patterns,
                )
                right_patterns = self._build_semantic_patterns(
                    source_role_bindings=right_roles,
                    enabled_patterns=enabled_patterns,
                )
                left_results = await self._generate_entity_candidate_chunk_results(
                    domain=domain,
                    business_summary=business_summary,
                    ontology_design_document=ontology_design_document,
                    rule_summary=rule_summary,
                    selected_table_schema={"table_count": len(left), "tables": left},
                    chunk=left,
                    enabled_patterns=enabled_patterns,
                    config_id=config_id,
                    table_roles=left_roles,
                    semantic_patterns=left_patterns,
                )
                right_results = await self._generate_entity_candidate_chunk_results(
                    domain=domain,
                    business_summary=business_summary,
                    ontology_design_document=ontology_design_document,
                    rule_summary=rule_summary,
                    selected_table_schema={"table_count": len(right), "tables": right},
                    chunk=right,
                    enabled_patterns=enabled_patterns,
                    config_id=config_id,
                    table_roles=right_roles,
                    semantic_patterns=right_patterns,
                )
                return left_results + right_results

            table = chunk[0]
            split_tables = self._bisect_single_table_segment(table)
            if len(split_tables) <= 1:
                raise
            merged_results: List[Dict[str, Any]] = []
            for split_table in split_tables:
                item_roles = [item for item in table_roles if (item.get("table_name") or "").upper() == (split_table.get("table_name") or "").upper()]
                item_patterns = self._build_semantic_patterns(
                    source_role_bindings=item_roles,
                    enabled_patterns=enabled_patterns,
                )
                merged_results.extend(
                    await self._generate_entity_candidate_chunk_results(
                        domain=domain,
                        business_summary=business_summary,
                        ontology_design_document=ontology_design_document,
                        rule_summary=rule_summary,
                        selected_table_schema={"table_count": 1, "tables": [split_table]},
                        chunk=[split_table],
                        enabled_patterns=enabled_patterns,
                        config_id=config_id,
                        table_roles=item_roles,
                        semantic_patterns=item_patterns,
                    )
                )
            return merged_results

    def _bisect_single_table_segment(self, table: Dict[str, Any]) -> List[Dict[str, Any]]:
        columns = list(table.get("columns") or [])
        if len(columns) <= 1:
            return [table]

        mid = len(columns) // 2
        first_columns = columns[:mid]
        second_columns = columns[mid:]
        base_start = int(table.get("segment_column_start") or 1)

        def build_segment(segment_columns: List[Dict[str, Any]], index: int, start: int) -> Dict[str, Any]:
            return {
                **table,
                "columns": segment_columns,
                "selected_column_count": len(segment_columns),
                "omitted_column_count": max(0, int(table.get("total_columns") or len(columns)) - len(segment_columns)),
                "segment_index": index,
                "segment_count": 2,
                "segment_column_start": start,
                "segment_column_end": start + len(segment_columns) - 1,
            }

        first = build_segment(first_columns, 1, base_start)
        second = build_segment(second_columns, 2, base_start + len(first_columns))
        logger.info(
            "Guide single table bisected after runtime reject: table=%s columns=%s split=(%s,%s)",
            table.get("table_name"),
            len(columns),
            len(first_columns),
            len(second_columns),
        )
        return [first, second]

    def _build_blueprint_table_details(self, table_details: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        blueprint_tables: List[Dict[str, Any]] = []
        for table in table_details:
            raw_columns = table.get("columns") or []
            compact_table = {
                "table_name": table.get("table_name"),
                "total_columns": len(raw_columns),
                "selected_column_count": len(raw_columns),
                "omitted_column_count": 0,
                "columns": [self._compact_column_for_llm(item) for item in raw_columns],
            }
            self._add_non_blank_value(compact_table, "owner", table.get("owner"))
            self._add_non_blank_value(compact_table, "table_comment", table.get("table_comment"))
            self._add_non_blank_value(compact_table, "source_role", table.get("source_role"))
            blueprint_tables.append(compact_table)
        return blueprint_tables

    def _parse_business_document_structured(self, business_document: str) -> Dict[str, Any]:
        normalized = self._normalize_document_text(business_document or "")
        lines = [line.strip() for line in normalized.splitlines() if line.strip()]
        keywords = []
        for token in ["SFR", "SPEC_LIMIT", "VCM_ID", "SensorID", "FTU", "FTD", "AA", "LBI", "socket", "carrier", "tooling", "Lens", "Lot", "供应商"]:
            if token.lower() in normalized.lower():
                keywords.append(token)
        focus_processes = [token for token in ["AA", "LBI", "FTU", "FTD"] if token.lower() in normalized.lower()]
        return {
            "raw_text": normalized,
            "document_outline": lines[:20],
            "keywords": keywords,
            "focus_processes": focus_processes,
            "focus_objects": [token for token in ["VCM_ID", "SensorID", "SocketID", "Lens", "Lot", "供应商"] if token.lower() in normalized.lower()],
            "focus_questions": [line for line in lines if "分析" in line or "根因" in line][:10],
        }

    def _extract_document_facts(
        self,
        business_document_parsed: Dict[str, Any],
        guide_context: Dict[str, Any],
        rule_analysis: Dict[str, Any],
        schema_analysis: Dict[str, Any],
    ) -> Dict[str, Any]:
        raw_text = str(business_document_parsed.get("raw_text") or "")
        upper_text = raw_text.upper()
        product_codes = self._extract_distinct_regex_tokens(upper_text, r"\b(?:PDX|MEM)\d+\b")
        history_sources = list(guide_context.get("history_case_sources") or [])
        if "FACA" in upper_text and "FACA" not in history_sources:
            history_sources.append("FACA")
        if not history_sources and any(token in upper_text for token in ["PDX", "PDX25", "历史", "CASE"]):
            history_sources.extend([token for token in ["PDX", "PDX25"] if token in upper_text])

        trace_chain = []
        if "缺陷" in raw_text or "SFR" in upper_text:
            trace_chain.append("缺陷现象")
        if any(token in upper_text for token in ["测试", "SFRMACRO", "SFRSUPERMACRO", "DARK-B", "DARK_B"]):
            trace_chain.append("测试项")
        if any(token in upper_text for token in ["BARCODE", "VCM_ID", "MODULE_ID"]):
            trace_chain.append("产品标识")
        if any(token in upper_text for token in ["站位", "LBI", "AA", "FTU", "FTD"]):
            trace_chain.append("过程站位")
        if any(token in upper_text for token in ["设备", "治具", "TOOLING", "CARRIER", "SOCKET"]):
            trace_chain.append("设备治具物料")
        if history_sources:
            trace_chain.append("历史案例")
        if any(token in raw_text for token in ["根因", "原因"]):
            trace_chain.append("根因方向")

        analysis_goals = []
        if any(token in raw_text for token in ["根因", "原因"]):
            analysis_goals.append("根因分析")
        if any(token in raw_text for token in ["追溯", "回溯"]):
            analysis_goals.append("多跳追溯")
        if any(token in raw_text for token in ["历史案例", "案例复用", "FACA"]):
            analysis_goals.append("历史案例复用")
        if any(token in raw_text for token in ["影响", "扩散", "范围"]):
            analysis_goals.append("影响分析")
        if any(token in raw_text for token in ["解释", "说明"]):
            analysis_goals.append("根因路径解释")

        business_scenario = guide_context.get("business_scenario")
        if not business_scenario:
            if "SFR" in upper_text and any(token in raw_text for token in ["根因", "缺陷"]):
                business_scenario = "DEFECT_ROOTCAUSE"
            elif "缺陷" in raw_text:
                business_scenario = "DEFECT_ROOTCAUSE"

        return {
            "scenario_name": "缺陷根因分析业务语义" if business_scenario == "DEFECT_ROOTCAUSE" else "业务分析场景",
            "business_scenario": business_scenario,
            "product_codes": product_codes,
            "history_knowledge_sources": history_sources,
            "trace_chain": trace_chain,
            "analysis_goals": analysis_goals,
            "focus_processes": business_document_parsed.get("focus_processes") or [],
            "focus_objects": business_document_parsed.get("focus_objects") or [],
            "document_keywords": business_document_parsed.get("keywords") or [],
            "document_outline": business_document_parsed.get("document_outline") or [],
            "rule_scope_hint": rule_analysis.get("primary_metric_families") or [],
            "key_tables_hint": schema_analysis.get("key_tables") or {},
        }

    def _parse_rule_datasets_from_text(self, text: str) -> List[Dict[str, Any]]:
        datasets_by_table: Dict[str, Dict[str, Any]] = {}
        insert_pattern = re.compile(
            r"INSERT\s+INTO\s+([^\s(]+)\s*\((.*?)\)\s*VALUES\s*\((.*?)\)\s*;",
            re.IGNORECASE | re.DOTALL,
        )
        for match in insert_pattern.finditer(text):
            full_table_name = (match.group(1) or "").strip()
            _owner, table_name = self._split_owner_and_table_name(full_table_name)
            columns = [
                self._normalize_sql_identifier(item)
                for item in self._split_sql_top_level(match.group(2) or "")
                if self._normalize_sql_identifier(item)
            ]
            raw_values = self._split_sql_top_level(match.group(3) or "")
            if not table_name or not columns or len(columns) != len(raw_values):
                continue

            record = {
                column_name: self._parse_sql_literal(raw_values[index])
                for index, column_name in enumerate(columns)
            }
            dataset = datasets_by_table.get(table_name)
            if not dataset:
                dataset = {
                    "table_name": table_name,
                    "columns": columns,
                    "records": [],
                }
                datasets_by_table[table_name] = dataset
            dataset["records"].append(record)

        normalized: List[Dict[str, Any]] = []
        for dataset in datasets_by_table.values():
            classified = self._classify_rule_dataset(
                table_name=dataset["table_name"],
                columns=dataset["columns"],
                records=dataset["records"],
            )
            if classified:
                normalized.append(classified)
        return normalized

    def _parse_sql_literal(self, value: str) -> Any:
        raw = str(value or "").strip()
        if not raw:
            return ""
        upper_raw = raw.upper()
        if upper_raw == "NULL":
            return None
        if raw.startswith("'") and raw.endswith("'") and len(raw) >= 2:
            return raw[1:-1].replace("''", "'")
        if re.fullmatch(r"-?\d+", raw):
            try:
                return int(raw)
            except Exception:
                return raw
        if re.fullmatch(r"-?\d+\.\d+", raw):
            try:
                return float(raw)
            except Exception:
                return raw
        return raw

    def _normalize_rule_datasets(self, datasets: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        normalized: List[Dict[str, Any]] = []
        seen_keys = set()
        for dataset in datasets:
            if not isinstance(dataset, dict):
                continue
            table_name = self._normalize_sql_identifier(dataset.get("table_name"))
            columns = [
                self._normalize_sql_identifier(item)
                for item in (dataset.get("columns") or [])
                if self._normalize_sql_identifier(item)
            ]
            raw_records = dataset.get("records") or []
            records: List[Dict[str, Any]] = []
            for record in raw_records[: self.RULE_DATA_MAX_RECORDS]:
                values = record.get("values") if isinstance(record, dict) and "values" in record else record
                if not isinstance(values, dict):
                    continue
                records.append({
                    self._normalize_sql_identifier(key): value
                    for key, value in values.items()
                    if self._normalize_sql_identifier(key)
                })
            classified = self._classify_rule_dataset(
                table_name=table_name,
                columns=columns,
                records=records,
                explicit_rule_type=dataset.get("rule_type"),
            )
            if not classified:
                continue
            key = (
                classified.get("rule_type"),
                classified.get("table_name"),
            )
            if key in seen_keys:
                continue
            seen_keys.add(key)
            normalized.append(classified)
        return normalized

    def _classify_rule_dataset(
        self,
        table_name: str,
        columns: List[str],
        records: List[Dict[str, Any]],
        explicit_rule_type: Optional[str] = None,
    ) -> Optional[Dict[str, Any]]:
        normalized_table_name = self._normalize_sql_identifier(table_name)
        normalized_columns = [
            self._normalize_sql_identifier(item)
            for item in columns
            if self._normalize_sql_identifier(item)
        ]
        column_set = set(normalized_columns)
        rule_type = (explicit_rule_type or "").strip().upper()
        is_spec_limit_like = (
            rule_type == "SPEC_LIMIT"
            or ("DB_NAME" in column_set and ("LSL" in column_set or "USL" in column_set))
            or ("SPEC" in normalized_table_name and "LIMIT" in normalized_table_name)
        )
        if not is_spec_limit_like:
            return None

        trimmed_records = records[: self.RULE_DATA_MAX_RECORDS]
        threshold_records = [
            record for record in trimmed_records
            if record.get("DB_NAME") and (record.get("LSL") is not None or record.get("USL") is not None)
        ]
        families = sorted({
            str(record.get("SPEC_FAMILY") or "").strip()
            for record in trimmed_records
            if str(record.get("SPEC_FAMILY") or "").strip()
        })
        metrics = [
            str(record.get("DB_NAME") or record.get("SPEC_METRIC") or "").strip()
            for record in threshold_records
            if str(record.get("DB_NAME") or record.get("SPEC_METRIC") or "").strip()
        ]
        summary = {
            "families": families[:20],
            "threshold_rule_count": len(threshold_records),
            "record_count": len(trimmed_records),
            "metric_examples": metrics[:20],
            "has_concrete_thresholds": bool(threshold_records),
        }
        return {
            "rule_type": "SPEC_LIMIT",
            "table_name": normalized_table_name or "SPEC_LIMIT",
            "record_count": len(trimmed_records),
            "columns": normalized_columns,
            "records": trimmed_records,
            "summary": summary,
        }

    def _load_database_rule_datasets(
        self,
        source_id: str,
        schema: Optional[str],
        relation_tables: List[str],
        explicit_rule_table_name: Optional[str],
        table_details: List[Dict[str, Any]],
    ) -> List[Dict[str, Any]]:
        candidate_tables = []
        if explicit_rule_table_name:
            candidate_tables.append(explicit_rule_table_name.strip().upper())
        for table_name in relation_tables:
            upper_name = (table_name or "").strip().upper()
            if any(token in upper_name for token in ["SPEC", "LIMIT", "RULE", "THRESHOLD"]):
                candidate_tables.append(upper_name)
        candidate_tables = list(dict.fromkeys([item for item in candidate_tables if item]))

        if not candidate_tables:
            return []

        existing_details = {
            (item.get("table_name") or "").upper(): item
            for item in table_details
            if item.get("table_name")
        }
        datasets: List[Dict[str, Any]] = []
        for table_name in candidate_tables:
            detail = existing_details.get(table_name)
            sample_rows = (detail or {}).get("sample_rows") or []
            if not detail or len(sample_rows) < min(self.DATABASE_RULE_SAMPLE_LIMIT, 20):
                try:
                    detail = self.source_service.get_remote_table_detail(
                        source_id=source_id,
                        table_name=table_name,
                        schema=schema,
                        sample_limit=self.DATABASE_RULE_SAMPLE_LIMIT,
                    )
                except Exception as exc:
                    if explicit_rule_table_name and table_name == explicit_rule_table_name.strip().upper():
                        raise ValueError(f"未能读取指定规则表 {table_name}: {str(exc)}") from exc
                    logger.warning("Load database rule table detail failed: table=%s error=%s", table_name, str(exc))
                    detail = existing_details.get(table_name) or detail
            dataset = self._classify_rule_dataset(
                table_name=detail.get("table_name") or table_name,
                columns=[column.get("column_name") for column in (detail.get("columns") or []) if column.get("column_name")],
                records=detail.get("sample_rows") or [],
            )
            if dataset:
                datasets.append(dataset)
        return datasets

    def _normalize_ddl_tables(self, ddl_tables: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        normalized_tables: List[Dict[str, Any]] = []
        for table in ddl_tables:
            if not isinstance(table, dict):
                continue
            table_name = (table.get("table_name") or "").strip().upper()
            if not table_name:
                continue
            columns = []
            for index, column in enumerate(table.get("columns") or [], start=1):
                if not isinstance(column, dict):
                    continue
                column_name = (column.get("column_name") or "").strip().upper()
                if not column_name:
                    continue
                columns.append({
                    "column_name": column_name,
                    "data_type": (column.get("data_type") or "").strip().upper(),
                    "nullable": (column.get("nullable") or "Y").strip().upper()[:1] or "Y",
                    "comments": (column.get("comments") or "").strip(),
                    "is_primary_key": (column.get("is_primary_key") or "N").strip().upper()[:1] or "N",
                    "column_id": int(column.get("column_id") or index),
                })
            normalized_tables.append({
                "owner": (table.get("owner") or "").strip().upper(),
                "table_name": table_name,
                "table_comment": (table.get("table_comment") or table.get("comments") or "").strip(),
                "columns": columns,
                "sample_rows": [],
            })
        return normalized_tables

    def _parse_ddl_tables(self, ddl_text: str) -> List[Dict[str, Any]]:
        table_comments = self._parse_ddl_table_comments(ddl_text)
        column_comments = self._parse_ddl_column_comments(ddl_text)
        tables: List[Dict[str, Any]] = []
        cursor = 0
        upper_text = ddl_text.upper()
        while True:
            match = re.search(r"\bCREATE\s+TABLE\b", upper_text[cursor:], re.IGNORECASE)
            if not match:
                break
            start = cursor + match.start()
            header_match = re.search(r"\bCREATE\s+TABLE\s+([^\s(]+)\s*\(", ddl_text[start:], re.IGNORECASE)
            if not header_match:
                cursor = start + len("CREATE TABLE")
                continue
            full_table_name = (header_match.group(1) or "").strip()
            open_paren = start + header_match.end() - 1
            close_paren = self._find_matching_parenthesis(ddl_text, open_paren)
            if close_paren == -1:
                cursor = open_paren + 1
                continue
            body = ddl_text[open_paren + 1:close_paren]
            owner, table_name = self._split_owner_and_table_name(full_table_name)
            columns = self._parse_ddl_columns(body)
            if columns:
                comment_key = f"{owner}.{table_name}" if owner else table_name
                tables.append({
                    "owner": owner,
                    "table_name": table_name,
                    "table_comment": table_comments.get(comment_key) or table_comments.get(table_name) or "",
                    "columns": [
                        {
                            **column,
                            "comments": column_comments.get((table_name, column.get("column_name") or ""), ""),
                        }
                        for column in columns
                    ],
                })
            cursor = close_paren + 1
        return tables

    def _parse_ddl_table_comments(self, ddl_text: str) -> Dict[str, str]:
        comments: Dict[str, str] = {}
        pattern = re.compile(
            r"COMMENT\s+ON\s+TABLE\s+([^\s]+)\s+IS\s+'((?:''|[^'])*)'",
            re.IGNORECASE | re.MULTILINE,
        )
        for match in pattern.finditer(ddl_text):
            full_name = (match.group(1) or "").strip()
            owner, table_name = self._split_owner_and_table_name(full_name)
            key = f"{owner}.{table_name}" if owner else table_name
            comments[key] = (match.group(2) or "").replace("''", "'").strip()
        return comments

    def _parse_ddl_column_comments(self, ddl_text: str) -> Dict[tuple[str, str], str]:
        comments: Dict[tuple[str, str], str] = {}
        pattern = re.compile(
            r"COMMENT\s+ON\s+COLUMN\s+([^. \t\r\n]+(?:\.[^. \t\r\n]+)?)\.([^. \t\r\n]+)\s+IS\s+'((?:''|[^'])*)'",
            re.IGNORECASE | re.MULTILINE,
        )
        for match in pattern.finditer(ddl_text):
            table_name = self._normalize_sql_identifier(match.group(1))
            column_name = self._normalize_sql_identifier(match.group(2))
            comments[(table_name, column_name)] = (match.group(3) or "").replace("''", "'").strip()
        return comments

    def _find_matching_parenthesis(self, text: str, open_index: int) -> int:
        depth = 0
        in_single_quote = False
        i = open_index
        while i < len(text):
            ch = text[i]
            if ch == "'" and (i == 0 or text[i - 1] != "\\"):
                if in_single_quote and i + 1 < len(text) and text[i + 1] == "'":
                    i += 2
                    continue
                in_single_quote = not in_single_quote
            elif not in_single_quote:
                if ch == "(":
                    depth += 1
                elif ch == ")":
                    depth -= 1
                    if depth == 0:
                        return i
            i += 1
        return -1

    def _parse_ddl_columns(self, body: str) -> List[Dict[str, Any]]:
        segments = self._split_sql_top_level(body)
        columns: List[Dict[str, Any]] = []
        pending_primary_keys: set[str] = set()
        for segment in segments:
            normalized = segment.strip()
            if not normalized:
                continue
            upper_segment = normalized.upper()
            if upper_segment.startswith("CONSTRAINT") or upper_segment.startswith("PRIMARY KEY") or upper_segment.startswith("UNIQUE") or upper_segment.startswith("FOREIGN KEY") or upper_segment.startswith("CHECK"):
                if "PRIMARY KEY" in upper_segment:
                    pk_match = re.search(r"PRIMARY\s+KEY\s*\((.*?)\)", normalized, re.IGNORECASE | re.DOTALL)
                    if pk_match:
                        pending_primary_keys.update(
                            self._normalize_sql_identifier(item)
                            for item in pk_match.group(1).split(",")
                            if self._normalize_sql_identifier(item)
                        )
                continue
            parts = normalized.split()
            if not parts:
                continue
            column_name = self._normalize_sql_identifier(parts[0])
            remainder = normalized[len(parts[0]):].strip()
            data_type_match = re.split(r"\bNOT\s+NULL\b|\bNULL\b|\bDEFAULT\b|\bCONSTRAINT\b|\bPRIMARY\s+KEY\b|\bUNIQUE\b|\bREFERENCES\b|\bCHECK\b", remainder, maxsplit=1, flags=re.IGNORECASE)
            data_type = (data_type_match[0] or "").strip().rstrip(",")
            columns.append({
                "column_name": column_name,
                "data_type": data_type.upper(),
                "nullable": "N" if re.search(r"\bNOT\s+NULL\b", remainder, re.IGNORECASE) else "Y",
                "is_primary_key": "Y" if re.search(r"\bPRIMARY\s+KEY\b", remainder, re.IGNORECASE) else "N",
                "column_id": len(columns) + 1,
            })
        if pending_primary_keys:
            for column in columns:
                if (column.get("column_name") or "") in pending_primary_keys:
                    column["is_primary_key"] = "Y"
        return columns

    def _split_sql_top_level(self, sql_body: str) -> List[str]:
        parts: List[str] = []
        current: List[str] = []
        depth = 0
        in_single_quote = False
        i = 0
        while i < len(sql_body):
            ch = sql_body[i]
            if ch == "'" and (i == 0 or sql_body[i - 1] != "\\"):
                if in_single_quote and i + 1 < len(sql_body) and sql_body[i + 1] == "'":
                    current.append(ch)
                    current.append(sql_body[i + 1])
                    i += 2
                    continue
                in_single_quote = not in_single_quote
            if not in_single_quote:
                if ch == "(":
                    depth += 1
                elif ch == ")":
                    depth = max(0, depth - 1)
                elif ch == "," and depth == 0:
                    parts.append("".join(current).strip())
                    current = []
                    i += 1
                    continue
            current.append(ch)
            i += 1
        tail = "".join(current).strip()
        if tail:
            parts.append(tail)
        return parts

    def _split_owner_and_table_name(self, full_name: str) -> tuple[str, str]:
        normalized = (full_name or "").strip()
        if "." in normalized:
            owner, table_name = normalized.split(".", 1)
            return self._normalize_sql_identifier(owner), self._normalize_sql_identifier(table_name)
        return "", self._normalize_sql_identifier(normalized)

    def _normalize_sql_identifier(self, identifier: Any) -> str:
        value = str(identifier or "").strip().strip('"').strip()
        return value.upper()

    def _build_selected_table_schema(
        self,
        table_details: List[Dict[str, Any]],
        for_llm: bool = False,
    ) -> Dict[str, Any]:
        tables = []
        for table in table_details:
            columns = table.get("columns") or []
            primary_keys = [
                col.get("column_name")
                for col in columns
                if str(col.get("is_primary_key") or "").upper() == "Y"
            ] or self._infer_primary_keys(columns)
            compact_columns = []
            for column in columns:
                if for_llm:
                    compact_columns.append(self._compact_column_for_llm(column))
                    continue

                compact_columns.append({
                    "column_name": column.get("column_name"),
                    "data_type": column.get("data_type"),
                    "nullable": column.get("nullable"),
                    "comments": column.get("comments") or "",
                })

            compact_table = {
                "table_name": table.get("table_name"),
                "column_count": len(columns),
                "columns": compact_columns,
            }
            if for_llm:
                self._add_non_blank_value(compact_table, "table_comment", table.get("table_comment"))
                self._add_non_blank_value(compact_table, "source_role", table.get("source_role"))
                if primary_keys:
                    compact_table["primary_keys"] = primary_keys
            else:
                compact_table.update({
                    "table_comment": table.get("table_comment") or "",
                    "source_role": table.get("source_role") or "",
                    "primary_keys": primary_keys,
                })
            tables.append(compact_table)
        return {"table_count": len(tables), "tables": tables}

    def _build_rule_summary(
        self,
        selected_table_schema: Dict[str, Any],
        business_document_parsed: Dict[str, Any],
        rule_datasets: Optional[List[Dict[str, Any]]] = None,
    ) -> Dict[str, Any]:
        tables = selected_table_schema.get("tables") or []
        raw_text = (business_document_parsed.get("raw_text") or "").upper()
        mentions_spec_limit = "SPEC_LIMIT" in raw_text
        spec_limit_negated = any(
            token in raw_text
            for token in [
                "不使用SPEC_LIMIT",
                "不依赖SPEC_LIMIT",
                "不依赖任何规则表",
                "不使用任何规则表",
                "NO SPEC_LIMIT",
            ]
        )
        matched_tables = [
            table for table in tables
            if "SPEC" in (table.get("table_name") or "").upper() and "LIMIT" in (table.get("table_name") or "").upper()
        ]
        structured_spec_dataset = next(
            (
                item for item in (rule_datasets or [])
                if (item.get("rule_type") or "").upper() == "SPEC_LIMIT"
                and bool((item.get("summary") or {}).get("has_concrete_thresholds"))
            ),
            None,
        )
        spec_limit_summary: Dict[str, Any]
        if structured_spec_dataset:
            summary_payload = structured_spec_dataset.get("summary") or {}
            records = structured_spec_dataset.get("records") or []
            threshold_records = [
                item for item in records
                if item.get("DB_NAME") and (item.get("LSL") is not None or item.get("USL") is not None)
            ]
            rule_examples = []
            for item in threshold_records[:10]:
                db_name = item.get("DB_NAME")
                family = item.get("SPEC_FAMILY") or item.get("SPEC_METRIC") or ""
                lsl = item.get("LSL")
                usl = item.get("USL")
                example = f"{family}:{db_name}[LSL={lsl if lsl is not None else '-'}, USL={usl if usl is not None else '-'}]"
                rule_examples.append(example)
            spec_limit_summary = {
                "matched": True,
                "table_name": structured_spec_dataset.get("table_name") or "SPEC_LIMIT",
                "rule_count": len(threshold_records),
                "families": summary_payload.get("families") or [],
                "metrics": summary_payload.get("metric_examples") or [],
                "key_fields": structured_spec_dataset.get("columns") or [],
                "rule_examples": rule_examples,
                "summary": (
                    f"识别到规则数据表 {structured_spec_dataset.get('table_name') or 'SPEC_LIMIT'}，"
                    f"已解析 {len(records)} 条规则记录，其中 {len(threshold_records)} 条带有明确上下限；"
                    f"缺陷识别范围覆盖规格族 {(' / '.join(summary_payload.get('families') or []) or '未标注族别')}，"
                    "可直接作为缺陷识别依据。"
                ),
                "has_concrete_rule_data": True,
                "rule_source_mode": "structured_data",
                "scope_summary": {
                    "families": summary_payload.get("families") or [],
                    "threshold_rule_count": len(threshold_records),
                    "record_count": len(records),
                    "metric_examples": summary_payload.get("metric_examples") or [],
                    "oos_logic": "measured_value < LSL 或 measured_value > USL 时识别为缺陷/异常",
                },
            }
        elif not matched_tables:
            if spec_limit_negated:
                summary_text = "业务文档明确说明当前场景不依赖 SPEC_LIMIT 或业务规则表。"
            elif mentions_spec_limit:
                summary_text = "业务文档提及了 SPEC_LIMIT 规则表，但当前所选业务关系表中未包含对应规则表。"
            else:
                summary_text = "当前未在已选业务关系表中识别到业务规则表。"
            spec_limit_summary = {
                "matched": False,
                "table_name": "",
                "rule_count": 0,
                "families": [],
                "metrics": [],
                "key_fields": [],
                "rule_examples": [],
                "summary": summary_text,
                "has_concrete_rule_data": False,
                "rule_source_mode": "none",
                "scope_summary": {},
            }
        else:
            target = matched_tables[0] if matched_tables else {}
            key_fields = [
                column.get("column_name")
                for column in (target.get("columns") or [])
                if any(token in (column.get("column_name") or "").upper() for token in ["SPEC", "LIMIT", "LSL", "USL", "RULE", "METRIC", "DB_NAME"])
            ]
            spec_limit_summary = {
                "matched": bool(target),
                "table_name": target.get("table_name") or "SPEC_LIMIT",
                "rule_count": len(key_fields),
                "families": ["SFR"] if "SFR" in (business_document_parsed.get("raw_text") or "").upper() else [],
                "metrics": key_fields[:50],
                "key_fields": key_fields,
                "rule_examples": key_fields[:10],
                "summary": "SPEC_LIMIT 表用于描述关注测试项目及其上下限阈值，超过阈值即判定为缺陷或相关异常。",
                "has_concrete_rule_data": False,
                "rule_source_mode": "structure_only",
                "scope_summary": {
                    "families": ["SFR"] if "SFR" in (business_document_parsed.get("raw_text") or "").upper() else [],
                    "threshold_rule_count": 0,
                    "record_count": 0,
                    "metric_examples": key_fields[:10],
                    "oos_logic": "仅识别到规则表结构，尚未解析到具体规则数据",
                },
            }

        return {
            "matched": spec_limit_summary.get("matched", False),
            "rule_type": "SPEC_LIMIT" if spec_limit_summary.get("matched") else "",
            "rule_table_name": spec_limit_summary.get("table_name") or "",
            "rule_count": spec_limit_summary.get("rule_count") or 0,
            "summary": spec_limit_summary.get("summary") or "当前未识别到业务规则表。",
            "families": spec_limit_summary.get("families") or [],
            "key_fields": spec_limit_summary.get("key_fields") or [],
            "rule_examples": spec_limit_summary.get("rule_examples") or [],
            "negated": spec_limit_negated,
            "has_concrete_rule_data": spec_limit_summary.get("has_concrete_rule_data", False),
            "rule_source_mode": spec_limit_summary.get("rule_source_mode") or "none",
            "scope_summary": spec_limit_summary.get("scope_summary") or {},
            "spec_limit_summary": spec_limit_summary,
        }

    def _analyze_spec_limit_rule_data(
        self,
        rule_datasets: List[Dict[str, Any]],
        rule_summary: Dict[str, Any],
        guide_context: Dict[str, Any],
    ) -> Dict[str, Any]:
        spec_dataset = next(
            (
                item for item in rule_datasets
                if (item.get("rule_type") or "").upper() == "SPEC_LIMIT"
            ),
            None,
        )
        if not spec_dataset:
            return {
                "rule_type": rule_summary.get("rule_type") or "",
                "rule_table_name": rule_summary.get("rule_table_name") or "",
                "has_concrete_thresholds": bool(rule_summary.get("has_concrete_rule_data")),
                "family_stats": [],
                "primary_metric_families": guide_context.get("focus_metric_families") or [],
                "extension_metric_families": [],
                "oos_logic": (rule_summary.get("scope_summary") or {}).get("oos_logic") or "",
                "summary": rule_summary.get("summary") or "",
            }

        records = spec_dataset.get("records") or []
        family_index: Dict[str, Dict[str, Any]] = {}
        for record in records:
            family = str(record.get("SPEC_FAMILY") or record.get("SPEC_METRIC") or "UNSPECIFIED").strip() or "UNSPECIFIED"
            metric_name = str(record.get("DB_NAME") or record.get("SPEC_METRIC") or "").strip()
            has_threshold = record.get("LSL") is not None or record.get("USL") is not None
            bucket = family_index.setdefault(family, {
                "family_name": family,
                "metric_count": 0,
                "threshold_metric_count": 0,
                "metric_examples": [],
            })
            bucket["metric_count"] += 1
            if has_threshold:
                bucket["threshold_metric_count"] += 1
            if metric_name and metric_name not in bucket["metric_examples"] and len(bucket["metric_examples"]) < 12:
                bucket["metric_examples"].append(metric_name)

        family_stats = sorted(
            family_index.values(),
            key=lambda item: (-int(item.get("threshold_metric_count") or 0), -int(item.get("metric_count") or 0), str(item.get("family_name") or "")),
        )
        focus_metric_families = guide_context.get("focus_metric_families") or []
        if focus_metric_families:
            primary_metric_families = [item for item in focus_metric_families if item in family_index]
        else:
            preferred = {"DARK-B", "SFRMACRO", "SFRSUPERMACRO"}
            primary_metric_families = [
                item.get("family_name")
                for item in family_stats
                if str(item.get("family_name") or "").upper() in preferred
            ]
            if not primary_metric_families:
                primary_metric_families = [
                    item.get("family_name")
                    for item in family_stats
                    if int(item.get("threshold_metric_count") or 0) > 0
                ][:3]
        extension_metric_families = [
            item.get("family_name")
            for item in family_stats
            if item.get("family_name") not in set(primary_metric_families)
        ]
        metric_name_examples = []
        for item in family_stats:
            for metric_name in item.get("metric_examples") or []:
                if metric_name not in metric_name_examples:
                    metric_name_examples.append(metric_name)
                if len(metric_name_examples) >= 30:
                    break
            if len(metric_name_examples) >= 30:
                break

        return {
            "rule_type": "SPEC_LIMIT",
            "rule_table_name": spec_dataset.get("table_name") or rule_summary.get("rule_table_name") or "",
            "has_concrete_thresholds": True,
            "family_stats": family_stats,
            "primary_metric_families": primary_metric_families,
            "extension_metric_families": extension_metric_families,
            "threshold_rule_count": sum(int(item.get("threshold_metric_count") or 0) for item in family_stats),
            "record_count": len(records),
            "oos_logic": "当测量值小于 LSL 或大于 USL 时判定为超差；任一关键测项超差即可识别为相关缺陷/异常。",
            "metric_name_examples": metric_name_examples,
            "summary": rule_summary.get("summary") or "",
        }

    def _build_business_summary(
        self,
        domain: SysDomain,
        business_document_parsed: Dict[str, Any],
        rule_summary: Dict[str, Any],
        table_roles: List[Dict[str, Any]],
    ) -> Dict[str, Any]:
        role_counts: Dict[str, int] = {}
        for item in table_roles:
            role = item.get("source_role") or "other"
            role_counts[role] = role_counts.get(role, 0) + 1
        raw_text = business_document_parsed.get("raw_text") or ""
        # 文档出现“缺陷”“SFR”等词并不代表应启用制造缺陷模板；只有
        # 用户显式选择的缺陷根因分析业务语义才允许构造该专用摘要。
        selected_semantic_code = (self._business_type_context.get("type_code") or "").upper()
        is_defect_scenario = (
            selected_semantic_code == "MANUFACTURING_DEFECT"
            and ("缺陷" in raw_text or "SFR" in raw_text.upper())
        )
        default_goal = (
            "围绕缺陷/异常分析场景构建可用于根因分析、影响分析与追溯分析的本体对象、属性和关系。"
            if is_defect_scenario
            else "围绕当前业务场景构建可用于分析、管理与图谱建模的本体对象、属性和关系。"
        )
        default_definition = (
            rule_summary.get("summary")
            or ("当关键测量指标或业务规则超出阈值时判定为缺陷或异常。" if is_defect_scenario else "根据业务文档和规则摘要识别关键业务对象及其关系。")
        )
        if rule_summary.get("negated"):
            default_definition = "当前场景不依赖专门规则表，主要根据业务事件、过程数据和对象关联构建本体。"
        return {
            "scenario_name": getattr(domain, "domain_name", "") or "业务本体分析",
            "core_goal": default_goal,
            "business_rule_definition": default_definition,
            "rule_source_mode": rule_summary.get("rule_source_mode") or "none",
            "rule_scope_summary": rule_summary.get("scope_summary") or {},
            "trace_path": [
                "业务主体",
                "业务事件或单据",
                "规则、指标或权益",
                "生命周期状态",
                "交易、服务或处置结果",
            ] if not is_defect_scenario else [
                "SFR NG项",
                "VCM_ID / SensorID",
                "FTU / FTD 测试结果",
                "AA / LBI 等上游站位",
                "机台 / socket / carrier / tooling",
                "Lens / VCM / Lot / 供应商",
                "根因候选",
            ],
            "analysis_dimensions": ["规则判定", "测量指标", "过程站位", "设备工装", "物料批次", "供应商"] if is_defect_scenario else ["业务主体", "业务事件", "指标规则", "生命周期状态", "交易或服务结果"],
            "focus_processes": business_document_parsed.get("focus_processes") or [],
            "focus_objects": business_document_parsed.get("focus_objects") or [],
            "rule_understanding": rule_summary.get("summary") or "",
            "table_role_overview": role_counts,
        }

    def _analyze_source_schema_keywords(
        self,
        selected_table_schema: Dict[str, Any],
        table_roles: List[Dict[str, Any]],
        source_role_bindings: List[Dict[str, Any]],
        business_document_parsed: Dict[str, Any],
    ) -> Dict[str, Any]:
        tables = selected_table_schema.get("tables") or []
        role_index = {
            (item.get("table_name") or "").upper(): item
            for item in table_roles
        }
        archetypes = [
            self._classify_table_archetypes(table, role_index.get((table.get("table_name") or "").upper()) or {})
            for table in tables
        ]
        product_index_tables = [item for item in archetypes if item.get("archetype") == "unit_index"]
        process_tables = [item for item in archetypes if item.get("archetype") == "process_wide_event"]
        test_tables = [item for item in archetypes if item.get("archetype") == "test_wide_result"]
        alarm_tables = [item for item in archetypes if item.get("archetype") == "alarm_event_source"]
        aa_tables = [item for item in archetypes if item.get("archetype") == "aa_feature_source"]
        rule_tables = [item for item in archetypes if item.get("archetype") == "rule_spec"]
        history_tables = [item for item in archetypes if item.get("archetype") == "history_case_source"]
        focus_stations = self._extract_focus_stations_from_process_table(
            selected_table_schema=selected_table_schema,
            preferred_processes=business_document_parsed.get("focus_processes") or [],
        )

        return {
            "table_archetypes": archetypes,
            "key_tables": {
                "product_index_table": product_index_tables[0].get("table_name") if product_index_tables else "",
                "process_table": process_tables[0].get("table_name") if process_tables else "",
                "rule_tables": [item.get("table_name") for item in rule_tables],
                "test_tables": [item.get("table_name") for item in test_tables],
                "aa_feature_tables": [item.get("table_name") for item in aa_tables],
                "alarm_tables": [item.get("table_name") for item in alarm_tables],
                "history_case_tables": [item.get("table_name") for item in history_tables],
            },
            "focus_stations": focus_stations,
            "source_role_bindings": source_role_bindings,
        }

    def _classify_table_archetypes(
        self,
        table: Dict[str, Any],
        table_role: Dict[str, Any],
    ) -> Dict[str, Any]:
        table_name = (table.get("table_name") or "").upper()
        columns = table.get("columns") or []
        column_names = [(item.get("column_name") or "").upper() for item in columns if item.get("column_name")]
        column_set = set(column_names)
        source_role = (table_role.get("source_role") or table.get("source_role") or "").strip().lower()
        matched_features: List[str] = []
        archetype = "other"
        reason = "未命中当前结构化识别规则"

        if "SPEC" in table_name and "LIMIT" in table_name:
            archetype = "rule_spec"
            reason = "表名命中 SPEC/LIMIT 规则表特征"
            matched_features.extend(["SPEC", "LIMIT"])
        elif "ALARM" in table_name:
            archetype = "alarm_event_source"
            reason = "表名命中 ALARM 告警表特征"
            matched_features.append("ALARM")
        elif any(token in table_name for token in ["ACTIVE_ALIGNMENT", "_AA_", " AA ", "AALOG"]) or ("FOCUS" in " ".join(column_names) and "SCAN" in " ".join(column_names)):
            archetype = "aa_feature_source"
            reason = "表名或字段命中 AA / 对焦特征"
            matched_features.extend([token for token in ["ACTIVE_ALIGNMENT", "FOCUS", "SCAN"] if token in table_name or token in " ".join(column_names)])
        elif any(token in table_name for token in ["CASE", "FACA", "RCA", "ROOT_CAUSE"]):
            archetype = "history_case_source"
            reason = "表名命中历史案例 / 根因知识特征"
            matched_features.extend([token for token in ["CASE", "FACA", "RCA"] if token in table_name])
        elif "PROCESS" in table_name or source_role == "process_history" or self._count_station_prefix_columns(column_names) >= 4:
            archetype = "process_wide_event"
            reason = "表名、角色或字段结构命中过程宽表特征"
            matched_features.extend([token for token in ["PROCESS", "MC_ID", "INPUT_TIME", "OUTPUT_TIME"] if token in table_name or any(token in name for name in column_names)])
        elif "UNIT" in table_name or (source_role == "entity_master" and any(token in column_set for token in ["VCM_ID", "MODULE_ID", "SENSOR_ID"])):
            archetype = "unit_index"
            reason = "表名或关键标识字段命中产品主索引特征"
            matched_features.extend([token for token in ["UNIT", "VCM_ID", "MODULE_ID", "SENSOR_ID"] if token in table_name or token in column_set])
        elif source_role == "measurement" or (
            "PASS_FAIL_DESCRIPTION" in column_set
            and "VCM_ID" in column_set
            and len(column_names) >= 10
        ) or any(token in table_name for token in ["SFR", "DARK", "FTU", "FTD", "TEST"]):
            archetype = "test_wide_result"
            reason = "表名、角色或 PASS_FAIL_DESCRIPTION 结构命中测试宽表特征"
            matched_features.extend([token for token in ["PASS_FAIL_DESCRIPTION", "VCM_ID", "SFR", "DARK", "FTU", "FTD"] if token in table_name or token in column_set])

        return {
            "table_name": table_name,
            "source_role": source_role or "",
            "archetype": archetype,
            "reason": reason,
            "matched_features": matched_features,
            "column_count": len(column_names),
            "primary_keys": table.get("primary_keys") or [],
        }

    def _extract_focus_stations_from_process_table(
        self,
        selected_table_schema: Dict[str, Any],
        preferred_processes: List[str],
    ) -> List[Dict[str, Any]]:
        process_tables = [
            table for table in (selected_table_schema.get("tables") or [])
            if "PROCESS" in (table.get("table_name") or "").upper() or (table.get("source_role") or "").lower() == "process_history"
        ]
        station_index: Dict[str, Dict[str, Any]] = {}
        suffixes = [
            "_MC_ID", "_OP_ID", "_LOT_ID", "_SUB_LOT_ID", "_INPUT_TIME", "_OUTPUT_TIME",
            "_START_TIME", "_END_TIME", "_LENS_ID", "_RAW_MATERIAL_LOT", "_CRR_ID",
            "_CAVITY_ID", "_TOOLING", "_SOCKETID", "_CARRIER", "_RECIPE", "_MATERIAL_LOT",
        ]
        for table in process_tables:
            for column in table.get("columns") or []:
                column_name = (column.get("column_name") or "").upper()
                for suffix in suffixes:
                    if not column_name.endswith(suffix):
                        continue
                    station_code = column_name[: -len(suffix)]
                    if not station_code or "_" not in station_code:
                        continue
                    bucket = station_index.setdefault(station_code, {
                        "station_code": station_code,
                        "evidence_column_count": 0,
                        "sample_columns": [],
                    })
                    bucket["evidence_column_count"] += 1
                    if column_name not in bucket["sample_columns"] and len(bucket["sample_columns"]) < 6:
                        bucket["sample_columns"].append(column_name)
                    break

        preferred_tokens = {str(item or "").strip().upper() for item in preferred_processes if str(item or "").strip()}
        recommended = []
        for item in station_index.values():
            station_code = item.get("station_code") or ""
            score = int(item.get("evidence_column_count") or 0)
            if any(token in station_code for token in ["LBI", "AA", "FTU", "FTD"]):
                score += 20
            if any(token and token in station_code for token in preferred_tokens):
                score += 10
            recommended.append({
                **item,
                "recommended": any(token in station_code for token in ["LBI", "AA", "FTU", "FTD"]),
                "score": score,
            })
        recommended.sort(key=lambda item: (-int(item.get("score") or 0), str(item.get("station_code") or "")))
        return recommended[:12]

    def _derive_metric_semantics(
        self,
        rule_analysis: Dict[str, Any],
        business_document_parsed: Dict[str, Any],
    ) -> Dict[str, Any]:
        metrics_by_family: Dict[str, List[str]] = {}
        for family in rule_analysis.get("family_stats") or []:
            metrics_by_family[str(family.get("family_name") or "")] = list(family.get("metric_examples") or [])

        categories = [
            ("中心解析力偏低", [r"CEN[_A-Z0-9]*AVG", r"CEN[_A-Z0-9]*MIN"]),
            ("边缘解析力偏低", [r"EDGE[_A-Z0-9]*MIN", r"30F_MIN", r"60F_MIN"]),
            ("左右边缘不对称", [r"LR_EDGE_DELTA", r"LEFT", r"RIGHT"]),
            ("上下边缘不对称", [r"TB_EDGE_DELTA", r"TOP", r"BOTTOM"]),
            ("高频/角度解析力不足", [r"30F_", r"60F_", r"SFR"]),
            ("焦点补偿异常", [r"FPDC_", r"FOCUS", r"COMPENS"]),
        ]
        semantic_categories = []
        for label, patterns in categories:
            matched_metrics: List[str] = []
            matched_families: List[str] = []
            for family_name, metrics in metrics_by_family.items():
                family_hit = False
                for metric_name in metrics:
                    upper_name = metric_name.upper()
                    if any(re.search(pattern, upper_name) for pattern in patterns):
                        if metric_name not in matched_metrics:
                            matched_metrics.append(metric_name)
                        family_hit = True
                if family_hit and family_name not in matched_families:
                    matched_families.append(family_name)
            if matched_metrics:
                semantic_categories.append({
                    "semantic_label": label,
                    "matched_families": matched_families,
                    "metric_examples": matched_metrics[:12],
                })

        return {
            "semantic_categories": semantic_categories,
            "document_focus_processes": business_document_parsed.get("focus_processes") or [],
        }

    def _build_focus_scope(
        self,
        guide_context: Dict[str, Any],
        document_facts: Dict[str, Any],
        rule_analysis: Dict[str, Any],
        schema_analysis: Dict[str, Any],
    ) -> Dict[str, Any]:
        families = guide_context.get("focus_metric_families") or rule_analysis.get("primary_metric_families") or []
        station_candidates = schema_analysis.get("focus_stations") or []
        recommended_station_codes = [
            item.get("station_code")
            for item in station_candidates
            if item.get("station_code") and (item.get("recommended") or len(station_candidates) <= 6)
        ]
        stations = guide_context.get("focus_stations") or recommended_station_codes[:6]
        history_case_sources = guide_context.get("history_case_sources") or document_facts.get("history_knowledge_sources") or []

        return {
            "business_scenario": document_facts.get("business_scenario"),
            "focus_metric_families": families,
            "focus_stations": stations,
            "history_case_sources": history_case_sources,
            "product_codes": document_facts.get("product_codes") or [],
            "analysis_goals": document_facts.get("analysis_goals") or [],
            "key_tables": schema_analysis.get("key_tables") or {},
        }

    def _build_table_roles(self, table_details: List[Dict[str, Any]], rule_summary: Dict[str, Any]) -> List[Dict[str, Any]]:
        roles = []
        rule_table_name = (rule_summary.get("rule_table_name") or "").upper()
        for table in table_details:
            table_name = (table.get("table_name") or "").upper()
            role = self._normalize_source_role(table.get("source_role"))
            reason = "用户显式指定角色"
            if not role:
                if rule_table_name and table_name == rule_table_name:
                    role = "rule_catalog"
                    reason = "识别为业务规则目录表"
                else:
                    role = self._infer_source_role(table_name, table.get("columns", []) or [])
                    reason = "根据表名和字段启发式推断"
            table_role = {
                "table_name": table_name,
                "source_role": role,
                "source_role_label": self.SOURCE_ROLE_CATALOG.get(role, self.SOURCE_ROLE_CATALOG["other"]),
                "reason": reason,
            }
            self._add_non_blank_value(table_role, "table_comment", table.get("table_comment"))
            roles.append(table_role)
        return roles

    def _infer_primary_keys(self, columns: List[Dict[str, Any]]) -> List[str]:
        explicit = [col.get("column_name") for col in columns if str(col.get("is_primary_key") or "").upper() == "Y"]
        if explicit:
            return explicit
        candidates = [col.get("column_name") for col in columns if (col.get("column_name") or "").upper() in {"ID", "VCM_ID", "SENSORID", "SENSOR_ID", "DB_NAME"}]
        return [item for item in candidates if item]

    def _chunk_blueprint_tables(
        self,
        domain: SysDomain,
        business_document: str,
        table_details: List[Dict[str, Any]],
        source_role_bindings: List[Dict[str, Any]],
        enabled_patterns: List[str],
        input_budget_tokens: Optional[int] = None,
    ) -> List[List[Dict[str, Any]]]:
        expanded_tables = self._expand_blueprint_tables_for_budget(
            domain=domain,
            business_document=business_document,
            table_details=table_details,
            source_role_bindings=source_role_bindings,
            enabled_patterns=enabled_patterns,
            input_budget_tokens=input_budget_tokens,
        )
        chunks: List[List[Dict[str, Any]]] = []
        current_chunk: List[Dict[str, Any]] = []
        current_budget = int(input_budget_tokens or 0)

        for table in expanded_tables:
            candidate_chunk = current_chunk + [table]
            candidate_bindings = self._filter_source_role_bindings(source_role_bindings, candidate_chunk)
            candidate_patterns = self._build_semantic_patterns(
                source_role_bindings=candidate_bindings,
                enabled_patterns=enabled_patterns,
            )
            estimated_tokens = self._estimate_blueprint_chunk_tokens(
                domain=domain,
                business_document=business_document,
                relation_tables=candidate_chunk,
                source_role_bindings=candidate_bindings,
                semantic_patterns=candidate_patterns,
            )
            current_columns = sum(int(item.get("total_columns") or 0) for item in candidate_chunk)
            should_split = bool(current_chunk) and (
                len(candidate_chunk) > self.GUIDE_BLUEPRINT_MAX_TABLES_PER_CHUNK
                or current_columns > self.GUIDE_BLUEPRINT_MAX_TOTAL_COLUMNS
                or (current_budget and estimated_tokens > current_budget)
            )
            if should_split:
                chunks.append(current_chunk)
                current_chunk = []
                candidate_chunk = [table]

            current_chunk = candidate_chunk

        if current_chunk:
            chunks.append(current_chunk)

        logger.info(
            "Guide ontology blueprint chunked: table_count=%s expanded_table_count=%s chunk_count=%s total_columns=%s input_budget_tokens=%s",
            len(table_details),
            len(expanded_tables),
            len(chunks),
            sum(int(item.get("total_columns") or 0) for item in table_details),
            current_budget,
        )
        return chunks

    def _expand_blueprint_tables_for_budget(
        self,
        domain: SysDomain,
        business_document: str,
        table_details: List[Dict[str, Any]],
        source_role_bindings: List[Dict[str, Any]],
        enabled_patterns: List[str],
        input_budget_tokens: Optional[int] = None,
    ) -> List[Dict[str, Any]]:
        budget = int(input_budget_tokens or 0)
        if budget <= 0:
            return table_details

        expanded: List[Dict[str, Any]] = []
        for table in table_details:
            candidate_bindings = self._filter_source_role_bindings(source_role_bindings, [table])
            candidate_patterns = self._build_semantic_patterns(
                source_role_bindings=candidate_bindings,
                enabled_patterns=enabled_patterns,
            )
            estimated_tokens = self._estimate_blueprint_chunk_tokens(
                domain=domain,
                business_document=business_document,
                relation_tables=[table],
                source_role_bindings=candidate_bindings,
                semantic_patterns=candidate_patterns,
            )
            if estimated_tokens <= budget:
                expanded.append(table)
                continue

            split_tables = self._split_single_table_for_budget(
                domain=domain,
                business_document=business_document,
                table=table,
                source_role_bindings=source_role_bindings,
                enabled_patterns=enabled_patterns,
                input_budget_tokens=budget,
            )
            expanded.extend(split_tables)

        return expanded

    def _split_single_table_for_budget(
        self,
        domain: SysDomain,
        business_document: str,
        table: Dict[str, Any],
        source_role_bindings: List[Dict[str, Any]],
        enabled_patterns: List[str],
        input_budget_tokens: int,
    ) -> List[Dict[str, Any]]:
        columns = list(table.get("columns") or [])
        if not columns:
            return [table]

        segments: List[Dict[str, Any]] = []
        current_columns: List[Dict[str, Any]] = []

        def build_table_segment(segment_columns: List[Dict[str, Any]]) -> Dict[str, Any]:
            return {
                **table,
                "columns": segment_columns,
                "selected_column_count": len(segment_columns),
                "omitted_column_count": max(0, len(columns) - len(segment_columns)),
            }

        for column in columns:
            candidate_columns = current_columns + [column]
            candidate_table = build_table_segment(candidate_columns)
            candidate_bindings = self._filter_source_role_bindings(source_role_bindings, [candidate_table])
            candidate_patterns = self._build_semantic_patterns(
                source_role_bindings=candidate_bindings,
                enabled_patterns=enabled_patterns,
            )
            estimated_tokens = self._estimate_blueprint_chunk_tokens(
                domain=domain,
                business_document=business_document,
                relation_tables=[candidate_table],
                source_role_bindings=candidate_bindings,
                semantic_patterns=candidate_patterns,
            )

            if current_columns and estimated_tokens > input_budget_tokens:
                segments.append(build_table_segment(current_columns))
                current_columns = [column]
            else:
                current_columns = candidate_columns

        if current_columns:
            segments.append(build_table_segment(current_columns))

        total_segments = max(1, len(segments))
        normalized_segments: List[Dict[str, Any]] = []
        for index, item in enumerate(segments, start=1):
            normalized_segments.append({
                **item,
                "segment_index": index,
                "segment_count": total_segments,
                "segment_column_start": sum(len(seg.get("columns") or []) for seg in segments[: index - 1]) + 1,
                "segment_column_end": sum(len(seg.get("columns") or []) for seg in segments[:index]),
            })

        logger.info(
            "Guide single table split for budget: table=%s total_columns=%s segment_count=%s input_budget_tokens=%s",
            table.get("table_name"),
            len(columns),
            total_segments,
            input_budget_tokens,
        )
        return normalized_segments

    def _estimate_blueprint_chunk_tokens(
        self,
        domain: SysDomain,
        business_document: str,
        relation_tables: List[Dict[str, Any]],
        source_role_bindings: List[Dict[str, Any]],
        semantic_patterns: List[Dict[str, Any]],
    ) -> int:
        payload = {
            "domain_name": getattr(domain, "domain_name", ""),
            "domain_desc": getattr(domain, "domain_desc", ""),
            "business_document": (business_document or "").strip()[:8000],
            "source_role_bindings": source_role_bindings,
            "semantic_patterns": semantic_patterns,
            "relation_tables": relation_tables,
        }
        return self.llm_service.estimate_structured_payload_tokens(payload) + 2200

    def _filter_source_role_bindings(
        self,
        source_role_bindings: List[Dict[str, Any]],
        chunk_tables: List[Dict[str, Any]],
    ) -> List[Dict[str, Any]]:
        table_names = {
            (item.get("table_name") or "").upper()
            for item in chunk_tables
            if item.get("table_name")
        }
        return [
            item for item in source_role_bindings
            if (item.get("table_name") or "").upper() in table_names
        ]

    def _merge_entity_candidate_results(self, candidate_results: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        merged: List[Dict[str, Any]] = []
        index: Dict[str, Dict[str, Any]] = {}
        for result in candidate_results:
            for entity in result.get("entity_candidates") or []:
                if not isinstance(entity, dict):
                    continue
                name = (entity.get("entityName") or "").strip()
                if not name:
                    continue
                key = name.lower()
                existing = index.get(key)
                if not existing:
                    copied = {
                        **entity,
                        "sourceHints": list(entity.get("sourceHints") or []),
                        "sourceRoles": list(entity.get("sourceRoles") or []),
                        "properties": list(entity.get("properties") or []),
                    }
                    index[key] = copied
                    merged.append(copied)
                    continue

                existing["candidateLevel"] = self._merge_candidate_level(existing.get("candidateLevel"), entity.get("candidateLevel"))
                existing_sources = {str(item).strip().upper() for item in existing.get("sourceHints") or []}
                for source_hint in entity.get("sourceHints") or []:
                    normalized_hint = str(source_hint).strip().upper()
                    if normalized_hint and normalized_hint not in existing_sources:
                        existing.setdefault("sourceHints", []).append(normalized_hint)
                        existing_sources.add(normalized_hint)
                existing_roles = {str(item).strip().lower() for item in existing.get("sourceRoles") or []}
                for role in entity.get("sourceRoles") or []:
                    normalized_role = str(role).strip().lower()
                    if normalized_role and normalized_role not in existing_roles:
                        existing.setdefault("sourceRoles", []).append(normalized_role)
                        existing_roles.add(normalized_role)
                existing_props = {(prop.get("propertyName") or "").strip().lower() for prop in existing.get("properties") or [] if isinstance(prop, dict)}
                for prop in entity.get("properties") or []:
                    prop_name = (prop.get("propertyName") or "").strip().lower()
                    if prop_name and prop_name not in existing_props:
                        existing.setdefault("properties", []).append(prop)
                        existing_props.add(prop_name)
        return merged

    def _merge_candidate_level(self, first: Any, second: Any) -> str:
        order = {"HIGH": 3, "MEDIUM": 2, "LOW": 1}
        first_level = str(first or "MEDIUM").strip().upper()
        second_level = str(second or "MEDIUM").strip().upper()
        if order.get(second_level, 0) > order.get(first_level, 0):
            return second_level
        return first_level if first_level in order else "MEDIUM"

    def _build_mandatory_master_entity_candidates(
        self,
        selected_table_schema: Dict[str, Any],
        table_roles: List[Dict[str, Any]],
    ) -> List[Dict[str, Any]]:
        """Return deterministic candidates for selected warehouse/store master tables."""
        role_by_table = {
            (item.get("table_name") or "").strip().upper(): (item.get("source_role") or "").strip().lower()
            for item in table_roles
            if isinstance(item, dict)
        }
        candidates: List[Dict[str, Any]] = []
        for table in selected_table_schema.get("tables") or []:
            if not isinstance(table, dict):
                continue
            table_name = (table.get("table_name") or "").strip().upper()
            if not table_name or not any(token in table_name for token in ("WAREHOUSE", "STORE", "RETAIL")):
                continue
            primary_keys = [str(item).strip().upper() for item in (table.get("primary_keys") or []) if str(item).strip()]
            if not primary_keys:
                continue
            columns = [item for item in (table.get("columns") or []) if isinstance(item, dict) and item.get("column_name")]
            column_by_name = {(item.get("column_name") or "").strip().upper(): item for item in columns}
            primary_key = primary_keys[0]
            if primary_key not in column_by_name:
                continue
            entity_name = self._master_entity_name_from_table(table_name)
            display_name = "仓库" if "WAREHOUSE" in table_name else ("零售门店" if "RETAIL" in table_name else "门店")
            selected_columns = [column_by_name[primary_key]]
            for column in columns:
                column_name = (column.get("column_name") or "").strip().upper()
                if column_name == primary_key or len(selected_columns) >= 8:
                    continue
                if any(token in column_name for token in ("NAME", "CODE", "TYPE", "ADDRESS", "REGION", "CITY", "STATUS")):
                    selected_columns.append(column)
            properties = [
                {
                    "propertyName": self._property_name_from_column(column.get("column_name") or ""),
                    "propertyDisplayName": (column.get("comments") or column.get("column_name") or "").strip(),
                    "propertyDesc": f"来自主数据表 {table_name} 的 {column.get('column_name')}",
                    "dataType": (column.get("data_type") or "VARCHAR2").strip().upper(),
                    "isPrimaryKey": "Y" if (column.get("column_name") or "").strip().upper() in primary_keys else "N",
                    "isNullable": (column.get("nullable") or "Y").strip().upper()[:1] or "Y",
                    "sourceTable": table_name,
                    "sourceColumn": (column.get("column_name") or "").strip().upper(),
                    "sourceDataType": (column.get("data_type") or "VARCHAR2").strip().upper(),
                    "mappingType": "DIRECT",
                    "formula": None,
                }
                for column in selected_columns
            ]
            candidates.append({
                "entityName": entity_name,
                "entityDisplayName": display_name,
                "entityDesc": f"由业务主数据表 {table_name} 直接映射的{display_name}实体。",
                "candidateLevel": "HIGH",
                "buildType": "TABLE",
                "sourceHints": [table_name],
                "sourceRoles": [role_by_table.get(table_name) or "entity_master"],
                "properties": properties,
                "mandatory_master_entity": True,
            })
        return candidates

    @staticmethod
    def _master_entity_name_from_table(table_name: str) -> str:
        ignored = {"DIM", "DIMENSION", "TB", "TBL", "MST", "MASTER"}
        tokens = [item for item in re.split(r"[^A-Z0-9]+", table_name.upper()) if item and item not in ignored]
        return "".join(token.title() for token in tokens) or "MasterData"

    @staticmethod
    def _property_name_from_column(column_name: str) -> str:
        tokens = [item.lower() for item in re.split(r"[^A-Za-z0-9]+", str(column_name or "")) if item]
        return "_".join(tokens) or "id"

    def _filter_entity_candidates_by_design_document(
        self,
        entity_candidates: List[Dict[str, Any]],
        ontology_design_document: Dict[str, Any],
    ) -> List[Dict[str, Any]]:
        included = ontology_design_document.get("included_entities") or []
        if not included:
            return entity_candidates
        allowed_names = {
            (item.get("entityName") or "").strip().lower()
            for item in included
            if isinstance(item, dict) and (item.get("entityName") or "").strip()
        }
        if not allowed_names:
            return entity_candidates
        filtered = [
            item for item in entity_candidates
            if item.get("mandatory_master_entity")
            or (item.get("entityName") or "").strip().lower() in allowed_names
        ]
        logger.info(
            "Guide entity candidates filtered by design document: before=%s after=%s allowed=%s",
            len(entity_candidates),
            len(filtered),
            sorted(allowed_names),
        )
        return filtered or entity_candidates

    def _filter_relation_candidates_by_design_document(
        self,
        relation_candidates: List[Dict[str, Any]],
        entity_candidates: List[Dict[str, Any]],
        ontology_design_document: Dict[str, Any],
    ) -> List[Dict[str, Any]]:
        allowed_entities = {
            (item.get("entityName") or "").strip().lower()
            for item in entity_candidates
            if (item.get("entityName") or "").strip()
        }
        filtered_by_entity = [
            item for item in relation_candidates
            if (item.get("sourceEntityName") or "").strip().lower() in allowed_entities
            and (item.get("targetEntityName") or "").strip().lower() in allowed_entities
        ]

        included = ontology_design_document.get("included_relations") or []
        allowed_relation_names = {
            (item.get("relationName") or "").strip()
            for item in included
            if isinstance(item, dict) and (item.get("relationName") or "").strip()
        }
        if allowed_relation_names:
            filtered = [
                item for item in filtered_by_entity
                if (item.get("relationName") or "").strip() in allowed_relation_names
            ]
        else:
            filtered = filtered_by_entity
        logger.info(
            "Guide relation candidates filtered by design document: before=%s entity_scoped=%s after=%s",
            len(relation_candidates),
            len(filtered_by_entity),
            len(filtered),
        )
        return filtered or filtered_by_entity

    def _build_ontology_design(
        self,
        entity_candidates: List[Dict[str, Any]],
        relation_candidates: List[Dict[str, Any]],
    ) -> Dict[str, Any]:
        entity_display_map = {
            (entity.get("entityName") or "").strip(): (entity.get("entityDisplayName") or entity.get("entityName") or "").strip()
            for entity in entity_candidates
            if (entity.get("entityName") or "").strip()
        }
        entities = []
        for entity in entity_candidates:
            entities.append({
                "entityName": entity.get("entityName"),
                "entityDisplayName": entity.get("entityDisplayName"),
                "entityDesc": entity.get("entityDesc"),
                "buildType": entity.get("buildType") or "TABLE",
                "sourceHints": entity.get("sourceHints") or [],
                "properties": entity.get("properties") or [],
            })
        relations = []
        for relation in relation_candidates:
            source_entity_name = relation.get("sourceEntityName")
            target_entity_name = relation.get("targetEntityName")
            relations.append({
                "sourceEntityName": source_entity_name,
                "targetEntityName": target_entity_name,
                "relationName": self._prefer_chinese_relation_name(
                    relation_name=relation.get("relationName"),
                    relation_desc=relation.get("relationDesc"),
                    source_entity_name=source_entity_name,
                    target_entity_name=target_entity_name,
                    entity_display_map=entity_display_map,
                ),
                "relationPredicate": relation.get("relationPredicate") or relation.get("relation_predicate") or "",
                "relationType": relation.get("relationType") or "ASSOCIATION",
                "relationDesc": relation.get("relationDesc") or "",
                "evidenceTables": relation.get("evidenceTables") or [],
            })
        return {"entities": entities, "relations": relations}

    def _prefer_chinese_relation_name(
        self,
        relation_name: Any,
        relation_desc: Any,
        source_entity_name: Any,
        target_entity_name: Any,
        entity_display_map: Dict[str, str],
    ) -> str:
        normalized_name = str(relation_name or "").strip()
        source_name = str(source_entity_name or "").strip()
        target_name = str(target_entity_name or "").strip()
        source_label = entity_display_map.get(source_name) or source_name
        target_label = entity_display_map.get(target_name) or target_name

        if any("\u4e00" <= ch <= "\u9fff" for ch in normalized_name):
            return self._compact_relation_name(normalized_name, source_label, target_label)

        desc_candidate = self._extract_relation_name_from_desc(relation_desc)
        if desc_candidate:
            return self._compact_relation_name(desc_candidate, source_label, target_label)

        if source_label and target_label:
            return "关联"
        return "关联"

    def _compact_relation_name(self, relation_name: str, source_label: str, target_label: str) -> str:
        """将“源实体 + 谓词 + 目标实体”收敛为图边可读的短谓词。"""
        compact = re.sub(r"[\s，,。；;：:、]", "", relation_name or "")
        for label in sorted({str(source_label or "").strip(), str(target_label or "").strip()}, key=len, reverse=True):
            if label:
                compact = compact.replace(label, "")
        compact = compact.strip("的与和及并")
        if compact:
            return compact[:12]
        # 无法去除实体名称时，仍控制边名称长度；详细语义保留在 relationDesc。
        return re.sub(r"[\s，,。；;：:、]", "", relation_name or "")[:12] or "关联"

    def _extract_relation_name_from_desc(self, relation_desc: Any) -> str:
        text = str(relation_desc or "").strip()
        if not text or not any("\u4e00" <= ch <= "\u9fff" for ch in text):
            return ""

        cleaned = " ".join(text.split())
        for prefix in [
            "该关系表示",
            "该关系描述",
            "该关系用于描述",
            "表示",
            "描述",
            "用于描述",
            "用于表示",
        ]:
            if cleaned.startswith(prefix):
                cleaned = cleaned[len(prefix):].strip(" ：:，,。；;")
                break

        for separator in ["。", "；", ";", "，", ",", "：", ":"]:
            if separator in cleaned:
                cleaned = cleaned.split(separator, 1)[0].strip()
                break

        if not cleaned:
            return ""
        if len(cleaned) > 24:
            return ""
        return cleaned

    def _annotate_data_enrichment_support(
        self,
        entities: List[Dict[str, Any]],
        relations: List[Dict[str, Any]],
        selected_tables: List[str],
    ) -> Dict[str, Any]:
        """Separate logical ontology design from evidence-backed mapping candidates."""
        selected_table_set = {
            str(item).strip().upper()
            for item in selected_tables or []
            if str(item).strip()
        }
        entity_status_by_name: Dict[str, str] = {}
        property_supported_count = 0
        for entity in entities or []:
            properties = entity.get("properties") or []
            supported_properties = 0
            for prop in properties:
                if not isinstance(prop, dict):
                    continue
                source_table = str(prop.get("sourceTable") or prop.get("source_table") or "").strip().upper()
                source_column = str(prop.get("sourceColumn") or prop.get("source_column") or "").strip().upper()
                is_supported = bool(source_table in selected_table_set and source_column)
                prop["data_support_status"] = "DATA_SUPPORTED" if is_supported else "LOGICAL_ONLY"
                prop["data_support_reason"] = (
                    f"已匹配来源字段 {source_table}.{source_column}"
                    if is_supported else "当前选定数据对象中未找到可确认的来源字段"
                )
                if is_supported:
                    supported_properties += 1
                    property_supported_count += 1

            source_hints = {
                str(item).strip().upper() for item in (entity.get("sourceHints") or []) if str(item).strip()
            }
            has_supported_hint = bool(source_hints & selected_table_set)
            status = "DATA_SUPPORTED" if supported_properties or has_supported_hint else "LOGICAL_ONLY"
            entity["data_support_status"] = status
            entity["data_support_reason"] = (
                f"已由 {supported_properties} 个属性来源或源表 {'、'.join(sorted(source_hints & selected_table_set)) or '-'} 支撑"
                if status == "DATA_SUPPORTED" else "仅来自逻辑设计，当前选定数据对象未提供实体或属性来源"
            )
            entity_status_by_name[(entity.get("entityName") or "").strip().lower()] = status

        relation_counts = {"DATA_SUPPORTED": 0, "DATA_CANDIDATE": 0, "LOGICAL_ONLY": 0}
        for relation in relations or []:
            evidence_tables = {
                str(item).strip().upper() for item in (relation.get("evidenceTables") or []) if str(item).strip()
            }
            source_table = str(relation.get("sourceTable") or "").strip().upper()
            target_table = str(relation.get("targetTable") or "").strip().upper()
            has_table_evidence = bool(evidence_tables & selected_table_set) or (
                source_table in selected_table_set and target_table in selected_table_set
            )
            source_status = entity_status_by_name.get((relation.get("sourceEntityName") or "").strip().lower(), "LOGICAL_ONLY")
            target_status = entity_status_by_name.get((relation.get("targetEntityName") or "").strip().lower(), "LOGICAL_ONLY")
            has_join = bool((relation.get("joinCondition") or "").strip() or (relation.get("edgeSql") or "").strip())
            if source_status == "DATA_SUPPORTED" and target_status == "DATA_SUPPORTED" and has_table_evidence and has_join:
                status = "DATA_SUPPORTED"
                reason = "两端实体、关系证据表和实际关联条件均已具备"
            elif source_status == "DATA_SUPPORTED" and target_status == "DATA_SUPPORTED" and has_table_evidence:
                status = "DATA_CANDIDATE"
                reason = "两端实体和关系证据表已具备，但仍需确认实际 Join 条件"
            else:
                status = "LOGICAL_ONLY"
                reason = "当前选定数据对象未同时支撑两端实体及关系依据"
            relation["data_support_status"] = status
            relation["data_support_reason"] = reason
            relation_counts[status] += 1

        entity_supported_count = sum(1 for item in entities or [] if item.get("data_support_status") == "DATA_SUPPORTED")
        return {
            "logical_entity_count": len(entities or []),
            "data_supported_entity_count": entity_supported_count,
            "logical_only_entity_count": len(entities or []) - entity_supported_count,
            "data_supported_property_count": property_supported_count,
            "data_supported_relation_count": relation_counts["DATA_SUPPORTED"],
            "data_candidate_relation_count": relation_counts["DATA_CANDIDATE"],
            "logical_only_relation_count": relation_counts["LOGICAL_ONLY"],
            "selected_tables": sorted(selected_table_set),
        }

    def _annotate_integrated_generation_support(self, result: Dict[str, Any]) -> None:
        """Mark an integrated document + data-table generation as mapping-ready.

        This mode is not a document-only logical design: the user has explicitly
        supplied the business document and the source data objects in the same
        run.  The produced ontology therefore defines the scope for the next
        data-mapping stage.  Exact source fields and joins are still verified by
        that stage, but entities and relations must not be excluded beforehand.
        """
        selected_tables = sorted({
            str(item).strip().upper()
            for item in (result.get("selected_tables") or [])
            if str(item).strip()
        })
        entities = result.get("entities") or []
        relations = result.get("relations") or []
        property_count = 0
        for entity in entities:
            entity["data_support_status"] = "DATA_SUPPORTED"
            entity["data_support_reason"] = "一体化生成已同时使用需求文档和选定业务数据表；将在数据映射阶段校验具体字段来源。"
            for prop in entity.get("properties") or []:
                if not isinstance(prop, dict):
                    continue
                prop["data_support_status"] = "DATA_SUPPORTED"
                prop["data_support_reason"] = "所属实体来自一体化生成的业务数据范围；待数据映射阶段确认具体来源字段。"
                property_count += 1
        for relation in relations:
            relation["data_support_status"] = "DATA_SUPPORTED"
            relation["data_support_reason"] = "一体化生成已同时使用需求文档和业务数据表；将在数据映射阶段验证实际 Join 或关系表。"

        result["entity_candidates"] = entities
        result["relation_candidates"] = relations
        result["mapping_design"] = self._build_mapping_design(
            entities=entities,
            relations=relations,
            source_role_bindings=result.get("source_role_bindings") or result.get("table_roles") or [],
            semantic_patterns=result.get("semantic_patterns") or [],
        )
        result["data_support"] = {
            "support_mode": "INTEGRATED_GENERATION",
            "support_basis": "需求文档 + 已选业务数据表",
            "logical_entity_count": len(entities),
            "data_supported_entity_count": len(entities),
            "logical_only_entity_count": 0,
            "data_supported_property_count": property_count,
            "data_supported_relation_count": len(relations),
            "data_candidate_relation_count": 0,
            "logical_only_relation_count": 0,
            "selected_tables": selected_tables,
        }

    def _merge_blueprint_results(self, blueprint_results: List[Dict[str, Any]]) -> Dict[str, Any]:
        merged_entities: List[Dict[str, Any]] = []
        entity_index: Dict[str, Dict[str, Any]] = {}
        merged_relations: List[Dict[str, Any]] = []
        relation_keys = set()
        generation_modes = []
        models = []

        for result in blueprint_results:
            if not isinstance(result, dict):
                continue
            if result.get("generation_mode"):
                generation_modes.append(result.get("generation_mode"))
            if result.get("model"):
                models.append(result.get("model"))

            for entity in result.get("entities") or []:
                if not isinstance(entity, dict):
                    continue
                entity_name = (entity.get("entityName") or "").strip()
                if not entity_name:
                    continue
                normalized = entity_name.lower()
                existing = entity_index.get(normalized)
                if not existing:
                    copied = {
                        **entity,
                        "sourceHints": list(entity.get("sourceHints") or []),
                        "properties": list(entity.get("properties") or []),
                    }
                    entity_index[normalized] = copied
                    merged_entities.append(copied)
                    continue

                existing_sources = {str(item).strip().upper() for item in existing.get("sourceHints") or []}
                for source_hint in entity.get("sourceHints") or []:
                    normalized_hint = str(source_hint).strip().upper()
                    if normalized_hint and normalized_hint not in existing_sources:
                        existing.setdefault("sourceHints", []).append(normalized_hint)
                        existing_sources.add(normalized_hint)

                existing_props = {
                    (prop.get("propertyName") or "").strip().lower()
                    for prop in existing.get("properties") or []
                    if isinstance(prop, dict) and prop.get("propertyName")
                }
                for prop in entity.get("properties") or []:
                    property_name = (prop.get("propertyName") or "").strip().lower()
                    if not property_name or property_name in existing_props:
                        continue
                    existing.setdefault("properties", []).append(prop)
                    existing_props.add(property_name)

            for relation in result.get("relations") or []:
                if not isinstance(relation, dict):
                    continue
                key = (
                    (relation.get("sourceEntityName") or "").strip().lower(),
                    (relation.get("targetEntityName") or "").strip().lower(),
                    (relation.get("relationName") or "").strip(),
                    (relation.get("relationType") or "").strip().upper(),
                )
                if not key[0] or not key[1] or key in relation_keys:
                    continue
                relation_keys.add(key)
                merged_relations.append({
                    **relation,
                    "evidenceTables": list(relation.get("evidenceTables") or []),
                })

        return {
            "entities": merged_entities,
            "relations": merged_relations,
            "generation_mode": "chunked_full_columns" if len(blueprint_results) > 1 else (generation_modes[0] if generation_modes else "llm"),
            "model": models[0] if models else None,
        }

    def _build_llm_table_details(self, table_details: List[Dict[str, Any]]) -> tuple[List[Dict[str, Any]], Dict[str, Any]]:
        table_count = max(1, len(table_details))
        per_table_budget = max(
            self.GUIDE_LLM_MIN_COLUMNS_PER_TABLE,
            min(
                self.GUIDE_LLM_MAX_COLUMNS_PER_TABLE,
                self.GUIDE_LLM_MAX_TOTAL_COLUMNS // table_count,
            ),
        )

        compact_tables: List[Dict[str, Any]] = []
        total_columns = 0
        selected_columns = 0

        for table in table_details:
            raw_columns = table.get("columns") or []
            total_columns += len(raw_columns)

            selected = self._select_columns_for_llm(raw_columns, per_table_budget)
            selected_columns += len(selected)
            selected_names = [item.get("column_name") for item in selected if item.get("column_name")]
            compact_rows = self._compact_sample_rows_for_llm(
                sample_rows=table.get("sample_rows") or [],
                selected_column_names=selected_names,
            )

            compact_table = {
                "table_name": table.get("table_name"),
                "total_columns": len(raw_columns),
                "selected_column_count": len(selected),
                "omitted_column_count": max(0, len(raw_columns) - len(selected)),
                "columns": selected,
            }
            self._add_non_blank_value(compact_table, "owner", table.get("owner"))
            self._add_non_blank_value(compact_table, "table_comment", table.get("table_comment"))
            self._add_non_blank_value(compact_table, "source_role", table.get("source_role"))
            if compact_rows:
                compact_table["sample_rows"] = compact_rows
                compact_table["sample_row_count"] = len(compact_rows)
            compact_tables.append(compact_table)

        summary = {
            "fetched_table_count": len(table_details),
            "llm_table_count": len(compact_tables),
            "max_columns_per_table": per_table_budget,
            "total_columns_before_compaction": total_columns,
            "total_columns_after_compaction": selected_columns,
            "omitted_columns": max(0, total_columns - selected_columns),
            "max_sample_rows_per_table": self.GUIDE_LLM_MAX_SAMPLE_ROWS_PER_TABLE,
            "max_sample_columns_per_row": self.GUIDE_LLM_MAX_SAMPLE_COLUMNS_PER_ROW,
        }
        logger.info("Guide LLM context compacted: %s", summary)
        return compact_tables, summary

    def _select_columns_for_llm(self, columns: List[Dict[str, Any]], limit: int) -> List[Dict[str, Any]]:
        if len(columns) <= limit:
            return [self._compact_column_for_llm(item) for item in columns]

        ranked = sorted(
            enumerate(columns),
            key=lambda item: (
                self._score_column_for_llm(item[1], item[0]),
                -(item[1].get("column_id") or item[0]),
            ),
            reverse=True,
        )
        selected_indexes = sorted(
            index for index, _column in ranked[: max(1, limit)]
        )
        return [self._compact_column_for_llm(columns[index]) for index in selected_indexes]

    def _score_column_for_llm(self, column: Dict[str, Any], index: int) -> int:
        name = (column.get("column_name") or "").upper()
        comment = (column.get("comments") or "").upper()
        data_type = (column.get("data_type") or "").upper()

        score = 0
        if index < 3:
            score += 40
        if name == "ID" or name.endswith("_ID") or name.endswith("ID"):
            score += 220
        if any(token in name for token in ["KEY", "CODE", "NO", "NUM", "SN"]):
            score += 150
        if any(token in name for token in ["NAME", "DESC", "TYPE", "STATUS", "FLAG", "RESULT", "LEVEL", "GRADE", "CLASS", "CATEGORY"]):
            score += 130
        if any(token in name for token in ["TIME", "DATE", "CREATE", "UPDATE", "MODIFY", "TS"]):
            score += 120
        if any(token in name for token in ["VALUE", "MEASURE", "METRIC", "SCORE", "LIMIT", "THRESHOLD", "SPEC", "RULE", "ALARM", "ERROR", "DEFECT", "CAUSE", "ACTION", "STEP", "ROUTE", "STATION", "LINE", "LOT", "BATCH", "MODEL", "PRODUCT", "ITEM", "DEVICE", "EQUIP", "MACHINE", "WO", "ORDER", "UNIT"]):
            score += 110
        if any(token in comment for token in ["主键", "唯一", "标识", "编码", "名称", "状态", "时间", "测量", "规则", "阈值", "缺陷", "原因", "动作", "工序", "设备", "产品", "批次"]):
            score += 60
        if any(token in data_type for token in ["CLOB", "BLOB", "NCLOB", "XMLTYPE", "LONG"]):
            score -= 120
        return score

    def _compact_column_for_llm(self, column: Dict[str, Any]) -> Dict[str, Any]:
        compact_column: Dict[str, Any] = {}
        self._add_non_blank_value(compact_column, "column_name", column.get("column_name"))
        self._add_non_blank_value(compact_column, "data_type", column.get("data_type"))
        self._add_non_blank_value(compact_column, "comments", column.get("comments"))
        return compact_column

    def _add_non_blank_value(
        self,
        target: Dict[str, Any],
        key: str,
        value: Any,
    ) -> None:
        if value is None:
            return
        if isinstance(value, str):
            value = value.strip()
            if not value:
                return
        target[key] = value

    def _compact_sample_rows_for_llm(
        self,
        sample_rows: List[Dict[str, Any]],
        selected_column_names: List[str],
    ) -> List[Dict[str, Any]]:
        allowed = set(selected_column_names[: self.GUIDE_LLM_MAX_SAMPLE_COLUMNS_PER_ROW])
        compact_rows: List[Dict[str, Any]] = []

        for row in sample_rows[: self.GUIDE_LLM_MAX_SAMPLE_ROWS_PER_TABLE]:
            compact_row: Dict[str, Any] = {}
            for column_name in selected_column_names:
                if column_name not in allowed or column_name not in row:
                    continue
                compact_row[column_name] = self._compact_sample_value(row.get(column_name))
            compact_rows.append(compact_row)
        return compact_rows

    def _compact_sample_value(self, value: Any) -> Any:
        if value is None or isinstance(value, (int, float, bool)):
            return value
        text = str(value)
        if len(text) <= 80:
            return text
        return f"{text[:77]}..."

    def _normalize_table_bindings(
        self,
        relation_tables: List[str],
        table_bindings: List[Dict[str, Any]],
    ) -> List[Dict[str, str]]:
        seen_tables = set()
        normalized: List[Dict[str, str]] = []

        for item in table_bindings:
            if not isinstance(item, dict):
                continue
            table_name = (item.get("table_name") or "").strip().upper()
            if not table_name or table_name in seen_tables:
                continue
            seen_tables.add(table_name)
            normalized.append({
                "table_name": table_name,
                "source_role": self._normalize_source_role(item.get("source_role")),
            })

        for table_name in relation_tables:
            normalized_name = (table_name or "").strip().upper()
            if not normalized_name or normalized_name in seen_tables:
                continue
            seen_tables.add(normalized_name)
            normalized.append({
                "table_name": normalized_name,
                "source_role": "",
            })

        return normalized

    def _normalize_source_role(self, value: Any) -> str:
        normalized = str(value or "").strip().lower()
        return normalized if normalized in self.SOURCE_ROLE_CATALOG else ""

    def _prioritize_table_bindings_for_guide(
        self,
        bindings: List[Dict[str, str]],
        business_document: str,
    ) -> List[Dict[str, str]]:
        raw_text = (business_document or "").upper()

        def score(item: Dict[str, str]) -> tuple[int, int, int]:
            table_name = (item.get("table_name") or "").upper()
            source_role = self._normalize_source_role(item.get("source_role"))
            mentions_table = 0 if table_name and table_name in raw_text else 1
            explicit_rule_table = 0 if source_role == "rule_catalog" else 1
            inferred_rule_table = 0 if ("SPEC" in table_name and "LIMIT" in table_name) else 1
            return (mentions_table, explicit_rule_table, inferred_rule_table)

        prioritized = sorted(
            enumerate(bindings),
            key=lambda pair: (*score(pair[1]), pair[0]),
        )
        return [item for _, item in prioritized]

    def _resolve_binding_role(self, table_name: str, bindings: List[Dict[str, str]]) -> str:
        normalized_table_name = (table_name or "").strip().upper()
        for item in bindings:
            if item.get("table_name") == normalized_table_name:
                explicit_role = self._normalize_source_role(item.get("source_role"))
                if explicit_role:
                    return explicit_role
        return ""

    def _infer_source_role(self, table_name: str, columns: List[Dict[str, Any]]) -> str:
        upper_name = (table_name or "").upper()
        upper_columns = " ".join((item.get("column_name") or "").upper() for item in columns[:80])
        text = f"{upper_name} {upper_columns}"

        if any(token in upper_name for token in ["SPEC", "LIMIT", "RULE", "THRESHOLD"]):
            return "rule_catalog"
        if any(token in upper_name for token in ["CASE", "FACA", "RCA", "KNOWLEDGE"]):
            return "case_library"
        if any(token in upper_name for token in ["PROCESS", "ROUTE", "TRACE", "HISTORY"]):
            return "process_history"
        if any(token in upper_name for token in ["ALARM", "EVENT", "LOG"]) and "PROCESS" not in upper_name:
            return "event_log"
        if any(token in upper_name for token in ["TEST", "SFR", "APS", "DARK", "MEASURE", "FTU", "FTD", "METRIC"]):
            return "measurement"
        if any(token in upper_name for token in ["UNIT", "MASTER", "ITEM", "MODEL", "PRODUCT"]) or "VCM_ID" in text:
            return "entity_master"
        return "other"

    def _build_source_role_bindings(self, table_details: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        bindings = []
        for table in table_details:
            source_role = self._normalize_source_role(table.get("source_role"))
            if not source_role:
                source_role = self._infer_source_role(table.get("table_name") or "", table.get("columns", []) or [])
            binding = {
                "table_name": (table.get("table_name") or "").upper(),
                "source_role": source_role,
                "source_role_label": self.SOURCE_ROLE_CATALOG.get(source_role, self.SOURCE_ROLE_CATALOG["other"]),
            }
            self._add_non_blank_value(binding, "table_comment", table.get("table_comment"))
            bindings.append(binding)
        return bindings

    def _build_semantic_patterns(
        self,
        source_role_bindings: List[Dict[str, Any]],
        enabled_patterns: List[str],
    ) -> List[Dict[str, Any]]:
        explicit_enabled = {
            (item or "").strip().lower()
            for item in enabled_patterns
            if (item or "").strip().lower() in self._semantic_pattern_catalog
        }
        patterns = []
        for pattern_code, config in self._semantic_pattern_catalog.items():
            enabled = pattern_code in explicit_enabled
            patterns.append({
                "pattern_code": pattern_code,
                "pattern_name": config.get("name"),
                "description": config.get("description"),
                "enabled": enabled,
                "business_type": self._business_type_context,
            })
        return patterns

    def _build_mapping_design(
        self,
        entities: List[Dict[str, Any]],
        relations: List[Dict[str, Any]],
        source_role_bindings: List[Dict[str, Any]],
        semantic_patterns: List[Dict[str, Any]],
    ) -> Dict[str, Any]:
        role_by_table = {item["table_name"]: item["source_role"] for item in source_role_bindings}
        derived_entity_names = {
            entity_name.lower()
            for pattern in semantic_patterns
            if pattern.get("enabled")
            for entity_name in (pattern.get("derived_entities") or [])
        }
        entity_mappings = []
        for entity in entities:
            source_hints = [str(item).strip().upper() for item in (entity.get("sourceHints") or []) if str(item).strip()]
            source_roles = sorted({role_by_table.get(table_name, "other") for table_name in source_hints})
            build_mode = "DERIVED_VIEW" if entity.get("entityName", "").lower() in derived_entity_names else "VIEW"
            entity_mappings.append({
                "entity_name": entity.get("entityName"),
                "source_hints": source_hints,
                "source_roles": source_roles,
                "recommended_build_mode": build_mode,
                "mapping_status": "PENDING",
                "data_support_status": entity.get("data_support_status") or "DATA_SUPPORTED",
                "data_support_reason": entity.get("data_support_reason") or "",
            })

        relation_mappings = [
            {
                "relation_name": relation.get("relationName"),
                "source_entity_name": relation.get("sourceEntityName"),
                "target_entity_name": relation.get("targetEntityName"),
                "relation_cardinality": relation.get("relationCardinality") or relation.get("relation_cardinality") or relation.get("cardinality"),
                "evidence_tables": relation.get("evidenceTables") or [],
                "mapping_status": "PENDING",
                "data_support_status": relation.get("data_support_status") or "DATA_SUPPORTED",
                "data_support_reason": relation.get("data_support_reason") or "",
            }
            for relation in relations
        ]

        return {
            "table_roles": source_role_bindings,
            "entity_mappings": entity_mappings,
            "relation_mappings": relation_mappings,
        }

    def _build_deployment_design(
        self,
        domain: SysDomain,
        entities: List[Dict[str, Any]],
        relations: List[Dict[str, Any]],
        source_role_bindings: List[Dict[str, Any]],
        semantic_patterns: List[Dict[str, Any]],
    ) -> Dict[str, Any]:
        prefix = self._build_deployment_prefix(domain)
        semantic_views = []
        seen_view_names = set()
        for binding in source_role_bindings:
            role = binding.get("source_role") or "other"
            view_name = f"VW_{prefix}_{self._sanitize_view_token(binding.get('table_name') or role)}"
            if view_name in seen_view_names:
                continue
            seen_view_names.add(view_name)
            semantic_views.append({
                "view_name": view_name,
                "view_kind": "source_wrap",
                "source_role": role,
                "source_tables": [binding.get("table_name")],
                "purpose": f"面向 {self.SOURCE_ROLE_CATALOG.get(role, '其他')} 的语义视图层",
                "deploy": False,
                "deploy_reason": "默认仅作为部署设计草案展示，未确认前不直接落库。",
                "sql": None,
            })

        for pattern in semantic_patterns:
            if not pattern.get("enabled"):
                continue
            for derived_name in pattern.get("derived_entities") or []:
                view_name = f"VW_{prefix}_{self._sanitize_view_token(derived_name)}"
                if view_name in seen_view_names:
                    continue
                seen_view_names.add(view_name)
                semantic_views.append({
                    "view_name": view_name,
                    "view_kind": "derived",
                    "source_role": "derived",
                    "source_tables": [],
                    "purpose": f"由语义模式 {pattern.get('pattern_name')} 派生的语义对象视图",
                    "deploy": False,
                    "deploy_reason": "派生语义视图需要补充明确SQL后再落库。",
                    "sql": None,
                })

        return {
            "semantic_views": semantic_views,
            "edge_views": [],
            "property_graph": {
                "graph_name": f"{prefix}_PG",
                "vertex_entities": [item.get("entityName") for item in entities if item.get("entityName")],
                "edge_relations": [item.get("relationName") for item in relations if item.get("relationName")],
                "note": "DDL 阶段应基于已确认的映射与语义视图，继续生成 Oracle 26ai Property Graph 创建语句。",
            },
        }

    def _build_deployment_design_from_view_plan(
        self,
        canonical_model: Dict[str, Any],
        view_plan: Dict[str, Any],
    ) -> Dict[str, Any]:
        standardized_views = []
        for item in view_plan.get("standardized_views") or []:
            standardized_views.append({
                "view_name": item.get("view_name"),
                "view_kind": item.get("view_kind") or "standardized",
                "source_role": "standardized",
                "source_tables": item.get("source_tables") or [],
                "purpose": item.get("purpose") or "",
                "deploy": bool(item.get("deploy")),
                "deploy_reason": item.get("deploy_reason") or "",
                "sql": item.get("sql"),
            })
        edge_views = []
        for item in view_plan.get("edge_views") or []:
            edge_views.append({
                "view_name": item.get("view_name"),
                "purpose": item.get("purpose") or "",
                "deploy": bool(item.get("deploy")),
                "deploy_reason": item.get("deploy_reason") or "",
                "source_tables": item.get("source_views") or item.get("source_tables") or [],
                "sql": item.get("sql"),
            })

        graph_layer = view_plan.get("graph_layer") or {}
        return {
            "semantic_views": standardized_views,
            "edge_views": edge_views,
            "property_graph": {
                "graph_name": "ONTOLOGY_PG",
                "vertex_entities": graph_layer.get("vertex_entities") or [item.get("entityName") for item in canonical_model.get("entities") or [] if item.get("entityName")],
                "edge_relations": graph_layer.get("edge_relations") or [item.get("relationName") for item in canonical_model.get("relations") or [] if item.get("relationName")],
                "note": "阶段 3 已基于领域模板固定标准化视图与图层骨架，后续 DDL 阶段补齐 SQL。",
            },
        }

    def _merge_canonical_model_enrichment(
        self,
        canonical_model: Dict[str, Any],
        enrichment: Dict[str, Any],
    ) -> Dict[str, Any]:
        base_entities = list(canonical_model.get("entities") or [])
        enriched_entities = {
            (item.get("entityName") or "").strip(): item
            for item in (enrichment.get("entities") or [])
            if (item.get("entityName") or "").strip()
        }
        merged_entities = []
        for entity in base_entities:
            entity_name = (entity.get("entityName") or "").strip()
            enriched = enriched_entities.get(entity_name) or {}
            base_property_index = {
                (prop.get("propertyName") or "").strip(): prop
                for prop in (entity.get("properties") or [])
                if (prop.get("propertyName") or "").strip()
            }
            enriched_property_index = {
                (prop.get("propertyName") or "").strip(): prop
                for prop in (enriched.get("properties") or [])
                if (prop.get("propertyName") or "").strip()
            }
            merged_properties = []
            for property_name, base_prop in base_property_index.items():
                candidate = {**base_prop, **(enriched_property_index.get(property_name) or {})}
                merged_properties.append(candidate)
            merged_entities.append({
                **entity,
                **{k: v for k, v in enriched.items() if k != "properties"},
                "entityName": entity_name,
                "properties": merged_properties,
            })

        base_relations = list(canonical_model.get("relations") or [])
        enriched_relations = {
            (
                (item.get("sourceEntityName") or "").strip(),
                (item.get("targetEntityName") or "").strip(),
                (item.get("relationName") or "").strip(),
                (item.get("relationType") or "").strip().upper(),
            ): item
            for item in (enrichment.get("relations") or [])
            if (item.get("sourceEntityName") or "").strip()
            and (item.get("targetEntityName") or "").strip()
            and (item.get("relationName") or "").strip()
        }
        merged_relations = []
        for relation in base_relations:
            key = (
                (relation.get("sourceEntityName") or "").strip(),
                (relation.get("targetEntityName") or "").strip(),
                (relation.get("relationName") or "").strip(),
                (relation.get("relationType") or "").strip().upper(),
            )
            enriched = enriched_relations.get(key) or {}
            merged = {**relation, **enriched}
            if relation.get("evidenceTables") and not merged.get("evidenceTables"):
                merged["evidenceTables"] = relation.get("evidenceTables")
            merged_relations.append(merged)

        return {
            **canonical_model,
            "entities": merged_entities,
            "relations": merged_relations,
        }

    def _merge_view_plan_enrichment(
        self,
        view_plan: Dict[str, Any],
        enrichment: Dict[str, Any],
    ) -> Dict[str, Any]:
        base_semantic_views = list(view_plan.get("standardized_views") or [])
        enriched_semantic_views = {
            (item.get("view_name") or "").strip().upper(): item
            for item in (enrichment.get("semantic_views") or [])
            if (item.get("view_name") or "").strip()
        }
        merged_semantic_views = []
        for item in base_semantic_views:
            key = (item.get("view_name") or "").strip().upper()
            merged = {**item, **(enriched_semantic_views.get(key) or {})}
            merged_semantic_views.append(merged)

        base_edge_views = list(view_plan.get("edge_views") or [])
        enriched_edge_views = {
            (item.get("view_name") or "").strip().upper(): item
            for item in (enrichment.get("edge_views") or [])
            if (item.get("view_name") or "").strip()
        }
        merged_edge_views = []
        for item in base_edge_views:
            key = (item.get("view_name") or "").strip().upper()
            merged = {**item, **(enriched_edge_views.get(key) or {})}
            merged_edge_views.append(merged)

        graph_layer = view_plan.get("graph_layer") or {}
        property_graph = enrichment.get("property_graph") or {}
        merged_graph_layer = {
            **graph_layer,
            "graph_name": property_graph.get("graph_name") or graph_layer.get("graph_name"),
            "vertex_entities": property_graph.get("vertex_entities") or graph_layer.get("vertex_entities") or [],
            "edge_relations": property_graph.get("edge_relations") or graph_layer.get("edge_relations") or [],
            "note": property_graph.get("note") or graph_layer.get("note") or "",
        }
        return {
            **view_plan,
            "standardized_views": merged_semantic_views,
            "edge_views": merged_edge_views,
            "graph_layer": merged_graph_layer,
        }

    def _build_deployment_prefix(self, domain: SysDomain) -> str:
        raw_name = (getattr(domain, "domain_name", "") or "").upper()
        token = "".join(ch if ch.isascii() and ch.isalnum() else "_" for ch in raw_name).strip("_")
        if token:
            return token[:24]
        fallback = (getattr(domain, "domain_id", "") or "DOMAIN").upper()
        return "".join(ch if ch.isascii() and ch.isalnum() else "_" for ch in fallback).strip("_")[:24] or "DOMAIN"

    def _view_suffix_for_role(self, role: str) -> str:
        mapping = {
            "entity_master": "MASTER_V",
            "process_history": "PROCESS_V",
            "measurement": "MEASUREMENT_V",
            "rule_catalog": "RULE_V",
            "case_library": "CASE_V",
            "event_log": "EVENT_V",
            "reference_data": "REFERENCE_V",
            "other": "SEMANTIC_V",
        }
        return mapping.get(role, "SEMANTIC_V")

    def _sanitize_view_token(self, value: str) -> str:
        token = "".join(ch if ch.isascii() and ch.isalnum() else "_" for ch in (value or "").upper()).strip("_")
        return token[:24] or "DERIVED_V"

    def _ensure_blueprint_storage(self) -> None:
        SysOntologyBlueprint.__table__.create(bind=self.db.bind, checkfirst=True)

    def _save_blueprint_package(
        self,
        domain_id: str,
        source_id: Optional[str],
        schema: Optional[str],
        payload: Dict[str, Any],
        created_by: str,
        status: str,
    ) -> Dict[str, Any]:
        self._ensure_blueprint_storage()
        latest = (
            self.db.query(SysOntologyBlueprint)
            .filter(SysOntologyBlueprint.domain_id == domain_id)
            .order_by(SysOntologyBlueprint.version_no.desc(), SysOntologyBlueprint.created_at.desc())
            .first()
        )
        version_no = (latest.version_no + 1) if latest else 1
        summary = {
            "entity_count": len(payload.get("entities") or []),
            "relation_count": len(payload.get("relations") or []),
            "source_table_count": len(payload.get("selected_tables") or []),
            "pattern_count": len([item for item in (payload.get("semantic_patterns") or []) if item.get("enabled")]),
            "status": status,
        }
        record = SysOntologyBlueprint(
            blueprint_id=generate_id("bp"),
            domain_id=domain_id,
            source_id=source_id,
            schema_name=schema,
            version_no=version_no,
            status=status,
            blueprint_json=json.dumps(payload, ensure_ascii=False, default=str),
            summary_json=json.dumps(summary, ensure_ascii=False, default=str),
            created_by=created_by,
        )
        self.db.add(record)
        self.db.commit()
        self.db.refresh(record)
        return {
            "blueprint_id": record.blueprint_id,
            "version_no": record.version_no,
        }

    def mark_blueprint_status(self, blueprint_id: str, status: str) -> None:
        if not blueprint_id:
            return
        normalized_status = (status or "").strip()
        if len(normalized_status) > 20:
            raise ValueError(f"蓝图状态长度不能超过 20 个字符：{normalized_status}")
        self._ensure_blueprint_storage()
        record = (
            self.db.query(SysOntologyBlueprint)
            .filter(SysOntologyBlueprint.blueprint_id == blueprint_id)
            .first()
        )
        if not record:
            return
        record.status = normalized_status
        record.updated_at = self._utcnow()
        self.db.commit()

    def apply_blueprint(
        self,
        domain_id: str,
        blueprint: Dict[str, Any],
        overwrite_existing: bool = False,
        logical_only: bool = False,
        created_by: str = "unknown",
    ) -> Dict[str, Any]:
        entities_payload = blueprint.get("entities") or []
        relations_payload = blueprint.get("relations") or []
        if not entities_payload:
            raise ValueError("当前没有可应用的实体建议")

        existing_entities = (
            self.db.query(SysOntologyEntity)
            .options(selectinload(SysOntologyEntity.properties))
            .filter(SysOntologyEntity.domain_id == domain_id)
            .order_by(SysOntologyEntity.created_at)
            .all()
        )
        entity_index = {(entity.entity_name or "").lower(): entity for entity in existing_entities}
        next_position_index = len(existing_entities)

        entity_result = {"created": 0, "updated": 0, "reused": 0}
        property_result = {"created": 0, "updated": 0}
        relation_result = {"created": 0, "updated": 0, "reused": 0, "skipped": 0}
        pending_property_refs: List[tuple[SysOntologyProperty, Optional[str]]] = []

        applied_entities: List[Dict[str, Any]] = []
        for entity_data in entities_payload:
            entity_name = (entity_data.get("entityName") or "").strip()
            if not entity_name:
                continue

            entity = entity_index.get(entity_name.lower())
            if entity:
                if overwrite_existing:
                    entity.entity_display_name = entity_data.get("entityDisplayName") or entity.entity_display_name
                    entity.entity_desc = entity_data.get("entityDesc") or entity.entity_desc
                    entity.build_type = entity_data.get("buildType") or entity.build_type
                    governance_info = entity_data.get("governanceInfo") or entity_data.get("governance_info") or {}
                    if "objectType" in entity_data or "object_type" in entity_data:
                        entity.object_type = self._normalize_object_type(
                            entity_data.get("objectType") or entity_data.get("object_type"),
                            entity_name,
                            entity_data.get("entityDesc") or entity.entity_desc or "",
                        )
                    if isinstance(governance_info, dict):
                        if any(key in governance_info for key in ["status", "governanceStatus"]):
                            entity.governance_status = (str(governance_info.get("status") or governance_info.get("governanceStatus") or "").strip().upper() or None)
                        if any(key in governance_info for key in ["dataOwner", "data_owner"]):
                            entity.data_owner = (str(governance_info.get("dataOwner") or governance_info.get("data_owner") or "").strip() or None)
                        if any(key in governance_info for key in ["securityLevel", "security_level"]):
                            entity.security_level = (str(governance_info.get("securityLevel") or governance_info.get("security_level") or "").strip().upper() or None)
                        if any(key in governance_info for key in ["updateFrequency", "update_frequency"]):
                            entity.update_frequency = (str(governance_info.get("updateFrequency") or governance_info.get("update_frequency") or "").strip().upper() or None)
                        if "tags" in governance_info:
                            entity.governance_tags_json = self._normalize_governance_tags(governance_info.get("tags"))
                    entity.updated_at = self._utcnow()
                    if entity.build_type == "VIEW":
                        entity.table_name = f"ONTO_NODE_{entity.entity_name.upper()}_V"
                    else:
                        entity.table_name = f"ONTO_NODE_{entity.entity_name.upper()}"
                    entity_result["updated"] += 1
                    entity_action = "updated"
                else:
                    entity_result["reused"] += 1
                    entity_action = "reused"
            else:
                build_type = (entity_data.get("buildType") or "TABLE").upper()
                governance_info = entity_data.get("governanceInfo") or entity_data.get("governance_info") or {}
                entity = SysOntologyEntity(
                    entity_id=generate_id("ent"),
                    domain_id=domain_id,
                    entity_name=entity_name,
                    entity_display_name=entity_data.get("entityDisplayName"),
                    entity_desc=entity_data.get("entityDesc"),
                    object_type=self._normalize_object_type(
                        entity_data.get("objectType") or entity_data.get("object_type"),
                        entity_name,
                        entity_data.get("entityDesc") or "",
                    ),
                    build_type=build_type,
                    table_name=f"ONTO_NODE_{entity_name.upper()}_V" if build_type == "VIEW" else f"ONTO_NODE_{entity_name.upper()}",
                    status="DRAFT",
                    governance_status=(str(governance_info.get("status") or governance_info.get("governanceStatus") or "").strip().upper() or None) if isinstance(governance_info, dict) else None,
                    data_owner=(str(governance_info.get("dataOwner") or governance_info.get("data_owner") or "").strip() or None) if isinstance(governance_info, dict) else None,
                    security_level=(str(governance_info.get("securityLevel") or governance_info.get("security_level") or "").strip().upper() or None) if isinstance(governance_info, dict) else None,
                    update_frequency=(str(governance_info.get("updateFrequency") or governance_info.get("update_frequency") or "").strip().upper() or None) if isinstance(governance_info, dict) else None,
                    governance_tags_json=self._normalize_governance_tags(governance_info.get("tags")) if isinstance(governance_info, dict) else self._normalize_governance_tags(None),
                    graph_position=json.dumps(self._build_graph_position(next_position_index)),
                    created_by=created_by,
                )
                next_position_index += 1
                self.db.add(entity)
                entity_index[entity_name.lower()] = entity
                entity_result["created"] += 1
                entity_action = "created"

            if not logical_only:
                self._upsert_entity_mapping_seed(entity, entity_data, overwrite_existing=overwrite_existing, created_by=created_by)
            existing_props = list(entity.properties) if getattr(entity, "properties", None) else []
            property_index = {(prop.property_name or "").lower(): prop for prop in existing_props}
            for prop_data in entity_data.get("properties", [])[:30]:
                property_name = (prop_data.get("propertyName") or "").strip().lower()
                if not property_name:
                    continue
                prop = property_index.get(property_name)
                created_new_property = False
                if prop:
                    if overwrite_existing:
                        prop.property_display_name = prop_data.get("propertyDisplayName") or prop.property_display_name
                        prop.property_desc = prop_data.get("propertyDesc") or prop.property_desc
                        prop.data_type = prop_data.get("dataType") or prop.data_type
                        prop.is_primary_key = "Y" if str(prop_data.get("isPrimaryKey") or "N").upper() == "Y" else "N"
                        prop.is_nullable = "N" if str(prop_data.get("isNullable") or "Y").upper() == "N" else "Y"
                        if "unit" in prop_data:
                            prop.unit = (str(prop_data.get("unit") or "").strip() or None)
                        if "valueConstraint" in prop_data or "value_constraint" in prop_data:
                            prop.value_constraint = (str(prop_data.get("valueConstraint") or prop_data.get("value_constraint") or "").strip() or None)
                        if any(key in prop_data for key in ["usage", "usageCodes", "usage_codes"]):
                            prop.usage_codes_json = self._normalize_usage_codes(prop_data.get("usage") or prop_data.get("usageCodes") or prop_data.get("usage_codes"))
                        if "isRequiredFilter" in prop_data or "is_required_filter" in prop_data:
                            prop.is_required_filter = self._normalize_flag(prop_data.get("isRequiredFilter") or prop_data.get("is_required_filter"), "N")
                        prop.updated_at = self._utcnow()
                        property_result["updated"] += 1
                else:
                    created_new_property = True
                    prop = SysOntologyProperty(
                        property_id=generate_id("prop"),
                        entity_id=entity.entity_id,
                        property_name=property_name,
                        property_display_name=prop_data.get("propertyDisplayName"),
                        data_type=prop_data.get("dataType") or "VARCHAR2",
                        is_primary_key="Y" if str(prop_data.get("isPrimaryKey") or "N").upper() == "Y" else "N",
                        is_nullable="N" if str(prop_data.get("isNullable") or "Y").upper() == "N" else "Y",
                        unit=(str(prop_data.get("unit") or "").strip() or None),
                        value_constraint=(str(prop_data.get("valueConstraint") or prop_data.get("value_constraint") or "").strip() or None),
                        usage_codes_json=self._normalize_usage_codes(prop_data.get("usage") or prop_data.get("usageCodes") or prop_data.get("usage_codes")),
                        is_required_filter=self._normalize_flag(prop_data.get("isRequiredFilter") or prop_data.get("is_required_filter"), "N"),
                        property_desc=prop_data.get("propertyDesc"),
                        order_num=len(property_index),
                    )
                    self.db.add(prop)
                    property_index[property_name] = prop
                    existing_props.append(prop)
                    property_result["created"] += 1
                has_ref_payload = any(key in prop_data for key in ["ref", "refProperty", "ref_property", "refPropertyId", "ref_property_id"])
                if has_ref_payload or created_new_property:
                    pending_property_refs.append((prop, self._extract_property_ref_token(
                        prop_data.get("ref")
                        or prop_data.get("refProperty")
                        or prop_data.get("ref_property")
                        or prop_data.get("refPropertyId")
                        or prop_data.get("ref_property_id")
                    )))
                if not logical_only:
                    self._upsert_property_mapping_seed(prop, prop_data, overwrite_existing=overwrite_existing, created_by=created_by)

            applied_entities.append({
                "entity_id": entity.entity_id,
                "entity_name": entity.entity_name,
                "action": entity_action,
            })

        all_entities = list(entity_index.values())
        entity_name_by_id = {entity.entity_id: entity.entity_name for entity in all_entities}
        all_properties = (
            self.db.query(SysOntologyProperty)
            .join(SysOntologyEntity, SysOntologyProperty.entity_id == SysOntologyEntity.entity_id)
            .filter(SysOntologyEntity.domain_id == domain_id)
            .all()
        )
        property_by_id = {prop.property_id: prop for prop in all_properties}
        property_by_path = {
            f"{entity_name_by_id.get(prop.entity_id, '').lower()}.{(prop.property_name or '').lower()}": prop
            for prop in all_properties
            if entity_name_by_id.get(prop.entity_id) and prop.property_name
        }
        for prop, ref_token in pending_property_refs:
            if not ref_token:
                prop.ref_property_id = None
                continue
            referenced = self._resolve_property_ref_token(ref_token, property_by_id, property_by_path)
            if referenced and referenced.property_id != prop.property_id:
                prop.ref_property_id = referenced.property_id
            else:
                prop.ref_property_id = None

        existing_relations = self.db.query(SysOntologyRelation).filter(SysOntologyRelation.domain_id == domain_id).all()
        relation_index = {
            (
                relation.source_entity_id,
                relation.target_entity_id,
                (relation.relation_name or "").strip().lower(),
            ): relation
            for relation in existing_relations
        }

        applied_relations: List[Dict[str, Any]] = []
        for relation_data in relations_payload:
            source_name = (relation_data.get("sourceEntityName") or "").strip().lower()
            target_name = (relation_data.get("targetEntityName") or "").strip().lower()
            relation_name = (relation_data.get("relationName") or "").strip()
            if not source_name or not target_name or not relation_name:
                relation_result["skipped"] += 1
                continue

            source_entity = entity_index.get(source_name)
            target_entity = entity_index.get(target_name)
            if not source_entity or not target_entity or source_entity.entity_id == target_entity.entity_id:
                relation_result["skipped"] += 1
                continue

            relation_type = (relation_data.get("relationType") or "ASSOCIATION").upper()
            if relation_type not in {"ONE_TO_ONE", "ONE_TO_MANY", "MANY_TO_MANY", "INHERITANCE", "ASSOCIATION"}:
                relation_type = "ASSOCIATION"
            relation_table_name = self._resolve_blueprint_relation_table_name(
                relation_data=relation_data,
                source_entity_name=source_entity.entity_name,
                target_entity_name=target_entity.entity_name,
                relation_type=relation_type,
            )

            relation_key = (
                source_entity.entity_id,
                target_entity.entity_id,
                relation_name.lower(),
            )
            relation = relation_index.get(relation_key)
            if relation:
                if overwrite_existing:
                    relation.relation_type = relation_type
                    relation.relation_desc = relation_data.get("relationDesc") or relation.relation_desc
                    relation.relation_table_name = relation_table_name
                    relation.updated_at = self._utcnow()
                    relation_result["updated"] += 1
                    relation_action = "updated"
                else:
                    if not (relation.relation_table_name or "").strip():
                        relation.relation_table_name = relation_table_name
                        relation.updated_at = self._utcnow()
                    relation_result["reused"] += 1
                    relation_action = "reused"
            else:
                relation = SysOntologyRelation(
                    relation_id=generate_id("rel"),
                    domain_id=domain_id,
                    source_entity_id=source_entity.entity_id,
                    target_entity_id=target_entity.entity_id,
                    relation_name=relation_name,
                    relation_type=relation_type,
                    relation_desc=relation_data.get("relationDesc"),
                    relation_table_name=relation_table_name,
                )
                self.db.add(relation)
                relation_index[relation_key] = relation
                relation_result["created"] += 1
                relation_action = "created"

            if not logical_only:
                self._upsert_relation_mapping_seed(relation, relation_data, entity_index, overwrite_existing=overwrite_existing, created_by=created_by)
            applied_relations.append({
                "relation_name": relation_name,
                "source_entity_name": source_entity.entity_name,
                "target_entity_name": target_entity.entity_name,
                "action": relation_action,
            })

        self.db.commit()
        return {
            "entities": entity_result,
            "properties": property_result,
            "relations": relation_result,
            "applied_entities": applied_entities,
            "applied_relations": applied_relations,
        }

    def _resolve_blueprint_relation_table_name(
        self,
        relation_data: Dict[str, Any],
        source_entity_name: str,
        target_entity_name: str,
        relation_type: str,
    ) -> Optional[str]:
        explicit = str(
            relation_data.get("relationTableName")
            or relation_data.get("relation_table_name")
            or ""
        ).strip().upper()
        if explicit:
            return explicit if explicit.startswith("ONTO_") else f"ONTO_EDGE_{explicit}"

        predicate = self._build_relation_predicate_token(
            str(relation_data.get("relationName") or "").strip(),
            source_entity_name=source_entity_name,
            target_entity_name=target_entity_name,
            llm_predicate=str(relation_data.get("relationPredicate") or relation_data.get("relation_predicate") or "").strip(),
        )
        return f"ONTO_EDGE_{source_entity_name.upper()}_{predicate}_{target_entity_name.upper()}"

    def _build_relation_predicate_token(
        self,
        relation_name: str,
        source_entity_name: str,
        target_entity_name: str,
        llm_predicate: str = "",
    ) -> str:
        mapping = {
            "属于": "BELONGS_TO",
            "有测试": "HAS_TEST_RUN",
            "产生": "GENERATES",
            "对照": "CHECKS_SPEC",
            "指向": "INDICATES",
            "经过": "PASSES_THROUGH",
            "发生于": "OCCURS_AT",
            "运行于": "RUNS_ON",
            "使用": "USES",
            "消耗": "CONSUMES",
            "关联报警": "HAS_ALARM",
            "具有AA特征": "HAS_AA_FEATURE",
            "相似于": "SIMILAR_TO",
            "支持": "SUPPORTS",
            "解决": "RESOLVES",
            "影响": "IMPACTS",
        }
        if relation_name in mapping:
            return mapping[relation_name]

        # The LLM produces a predicate only (for example USES), never a full
        # source/target expression.  This avoids names such as
        # EQUIPMENT_REL_EQUIPMENT_WORKSHOP_WORKSHOP.
        predicate = re.sub(r"[^A-Z0-9]+", "_", llm_predicate.upper()).strip("_")
        if predicate and not predicate.startswith("ONTO_"):
            return predicate[:48]

        # A valid graph edge still needs a portable name when no configured
        # vocabulary or model suggestion is available.
        return "RELATED_TO"

    def _upsert_entity_mapping_seed(
        self,
        entity: SysOntologyEntity,
        entity_data: Dict[str, Any],
        *,
        overwrite_existing: bool,
        created_by: str,
    ) -> None:
        mapping = self.db.query(SysEntityMapping).filter(SysEntityMapping.entity_id == entity.entity_id).first()
        build_type = (entity_data.get("buildType") or entity.build_type or "TABLE").upper()
        if not mapping:
            mapping = SysEntityMapping(
                mapping_id=generate_id("emap"),
                entity_id=entity.entity_id,
                build_type=build_type,
                mapping_status="PENDING",
                mapped_by=created_by,
                mapped_at=self._utcnow(),
            )
            self.db.add(mapping)
            return
        if overwrite_existing or not (mapping.build_type or "").strip():
            mapping.build_type = build_type
        mapping.mapped_by = created_by
        mapping.mapped_at = self._utcnow()

    def _upsert_property_mapping_seed(
        self,
        prop: SysOntologyProperty,
        prop_data: Dict[str, Any],
        *,
        overwrite_existing: bool,
        created_by: str,
    ) -> None:
        source_table = (prop_data.get("sourceTable") or prop_data.get("source_table") or "").strip().upper()
        source_column = (prop_data.get("sourceColumn") or prop_data.get("source_column") or "").strip().upper()
        mapping_type = (prop_data.get("mappingType") or prop_data.get("mapping_type") or "DIRECT").strip().upper()
        formula = (prop_data.get("formula") or prop_data.get("formula_expr") or "").strip()
        source_data_type = (prop_data.get("sourceDataType") or prop_data.get("source_data_type") or "").strip().upper()
        if not source_table or not source_column:
            return

        mapping = self.db.query(SysPropertyMapping).filter(SysPropertyMapping.property_id == prop.property_id).first()
        if not mapping:
            mapping = SysPropertyMapping(
                mapping_id=generate_id("pmap"),
                property_id=prop.property_id,
                source_table=source_table,
                source_column=source_column,
                mapping_type=mapping_type,
                formula_expr=formula or None,
                formula_desc=(prop_data.get("propertyDesc") or prop_data.get("property_desc") or "").strip() or None,
                confidence="MEDIUM",
                mapping_status="SUGGESTED",
                mapped_by=created_by,
                mapped_at=self._utcnow(),
            )
            self.db.add(mapping)
        else:
            should_fill = overwrite_existing or not ((mapping.source_table or "").strip() and (mapping.source_column or "").strip())
            if should_fill:
                mapping.source_table = source_table
                mapping.source_column = source_column
                mapping.mapping_type = mapping_type
                mapping.formula_expr = formula or None
                if not (mapping.formula_desc or "").strip():
                    mapping.formula_desc = (prop_data.get("propertyDesc") or prop_data.get("property_desc") or "").strip() or None
                if mapping.mapping_status != "CONFIRMED":
                    mapping.mapping_status = "SUGGESTED"
            mapping.mapped_by = created_by
            mapping.mapped_at = self._utcnow()

        if source_data_type and not (prop.data_type or "").strip():
            prop.data_type = source_data_type

    def _upsert_relation_mapping_seed(
        self,
        relation: SysOntologyRelation,
        relation_data: Dict[str, Any],
        entity_index: Dict[str, SysOntologyEntity],
        *,
        overwrite_existing: bool,
        created_by: str,
    ) -> None:
        evidence_tables = [
            str(item).strip().upper()
            for item in (relation_data.get("evidenceTables") or relation_data.get("evidence_tables") or [])
            if str(item).strip()
        ]
        source_name = (relation_data.get("sourceEntityName") or "").strip().lower()
        target_name = (relation_data.get("targetEntityName") or "").strip().lower()
        source_entity = entity_index.get(source_name)
        target_entity = entity_index.get(target_name)
        source_table = (relation_data.get("sourceTable") or relation_data.get("source_table") or "").strip().upper()
        target_table = (relation_data.get("targetTable") or relation_data.get("target_table") or "").strip().upper()
        raw_relation_cardinality = (
            relation_data.get("relationCardinality")
            or relation_data.get("relation_cardinality")
            or relation_data.get("cardinality")
            or relation.relation_type
        )
        relation_cardinality = str(raw_relation_cardinality or "").strip().upper().replace("-", "_").replace(" ", "_") or None
        if relation_cardinality not in {"ONE_TO_ONE", "ONE_TO_MANY", "MANY_TO_ONE", "MANY_TO_MANY"}:
            relation_cardinality = None
        join_condition = (relation_data.get("joinCondition") or relation_data.get("join_condition") or "").strip() or None
        edge_sql = (relation_data.get("edgeSql") or relation_data.get("edge_sql") or "").strip() or None
        if not source_table:
            source_table = evidence_tables[0] if evidence_tables else ""
        if not target_table:
            target_table = evidence_tables[1] if len(evidence_tables) > 1 else (source_table or "")

        if not source_table and not target_table:
            return

        mapping = self.db.query(SysRelationMapping).filter(SysRelationMapping.relation_id == relation.relation_id).first()
        if not mapping:
            mapping = SysRelationMapping(
                mapping_id=generate_id("rmap"),
                relation_id=relation.relation_id,
                source_table=source_table or None,
                target_table=target_table or None,
                relation_cardinality=relation_cardinality,
                join_condition=join_condition,
                edge_sql=edge_sql,
                mapping_status="CONFIRMED" if edge_sql else "SUGGESTED",
                mapped_by=created_by,
                mapped_at=self._utcnow(),
            )
            self.db.add(mapping)
            return

        should_fill = overwrite_existing or not ((mapping.source_table or "").strip() or (mapping.target_table or "").strip() or (mapping.edge_sql or "").strip())
        if should_fill and mapping.mapping_status != "CONFIRMED":
            mapping.source_table = source_table or mapping.source_table
            mapping.target_table = target_table or mapping.target_table
            mapping.relation_cardinality = relation_cardinality or mapping.relation_cardinality
            mapping.join_condition = join_condition or mapping.join_condition
            mapping.edge_sql = edge_sql or mapping.edge_sql
            mapping.mapping_status = "CONFIRMED" if (mapping.edge_sql or "").strip() else "SUGGESTED"
        mapping.mapped_by = created_by
        mapping.mapped_at = self._utcnow()

    def _build_graph_position(self, index: int) -> Dict[str, int]:
        column = index % 4
        row = index // 4
        return {
            "x": 80 + column * 220,
            "y": 60 + row * 160,
        }

    def _utcnow(self):
        from datetime import datetime

        return datetime.utcnow()

    def _decode_text_bytes(self, content: bytes) -> str:
        for encoding in ("utf-8", "utf-8-sig", "gbk", "gb18030"):
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                continue
        return content.decode("latin-1", errors="ignore")

    def _extract_docx_markdown(self, content: bytes) -> str:
        """Convert DOCX to lightweight Markdown while retaining document structure.

        The earlier XML-only reader flattened headings, list hierarchy and tables into
        plain text. Those structures are meaningful inputs to ontology design, so use
        python-docx here and keep them as Markdown for the Guide and LLM.
        """
        try:
            from docx import Document
        except ImportError as exc:
            raise ValueError("当前运行环境未安装 python-docx，暂时无法解析 DOCX 文档") from exc

        try:
            document = Document(BytesIO(content))
        except Exception as exc:
            raise ValueError(f"无法解析 docx 文档: {str(exc)}") from exc

        parts: List[str] = []
        for paragraph in document.paragraphs:
            text = (paragraph.text or "").strip()
            if not text:
                continue
            style_name = (getattr(paragraph.style, "name", "") or "").lower()
            heading_match = re.search(r"heading\s*([1-6])|标题\s*([1-6])", style_name, re.IGNORECASE)
            if heading_match:
                level = next((int(value) for value in heading_match.groups() if value), 1)
                parts.append(f"{'#' * level} {text}")
            elif "list bullet" in style_name or "项目符号" in style_name:
                parts.append(f"- {text}")
            elif "list number" in style_name or "编号" in style_name:
                parts.append(f"1. {text}")
            else:
                parts.append(text)

        for table in document.tables:
            rows = [
                [self._escape_markdown_table_cell(cell.text) for cell in row.cells]
                for row in table.rows
            ]
            rows = [row for row in rows if any(cell.strip() for cell in row)]
            if not rows:
                continue
            column_count = max(len(row) for row in rows)
            normalized_rows = [row + [""] * (column_count - len(row)) for row in rows]
            parts.append("| " + " | ".join(normalized_rows[0]) + " |")
            parts.append("| " + " | ".join(["---"] * column_count) + " |")
            parts.extend("| " + " | ".join(row) + " |" for row in normalized_rows[1:])

        return "\n\n".join(parts)

    def _extract_docx_text(self, content: bytes) -> str:
        """Compatibility wrapper for callers that only need a text representation."""
        return self._extract_docx_markdown(content)

    @staticmethod
    def _escape_markdown_table_cell(value: str) -> str:
        return re.sub(r"\s+", " ", (value or "").replace("|", "\\|")).strip()

    def _text_to_markdown(self, text: str, preserve_markdown: bool = False) -> str:
        normalized = self._normalize_document_text(text)
        if not normalized:
            return ""
        if preserve_markdown:
            return normalized

        markdown_lines: List[str] = []
        for line in normalized.splitlines():
            if re.match(r"^(第[一二三四五六七八九十\d]+[章节]|[一二三四五六七八九十]+、|\d+(?:\.\d+)*[、.]?)\s*", line):
                markdown_lines.append(f"## {line}")
            elif re.match(r"^[•·▪◦]\s*", line):
                markdown_lines.append("- " + re.sub(r"^[•·▪◦]\s*", "", line))
            else:
                markdown_lines.append(line)
        return "\n\n".join(markdown_lines)

    @staticmethod
    def _normalize_markdown(text: str) -> str:
        lines = [line.rstrip() for line in (text or "").replace("\r\n", "\n").split("\n")]
        while lines and not lines[0].strip():
            lines.pop(0)
        while lines and not lines[-1].strip():
            lines.pop()
        return "\n".join(lines).strip()

    def _extract_docx_text_legacy(self, content: bytes) -> str:
        namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        parts: List[str] = []
        try:
            with ZipFile(BytesIO(content)) as archive:
                with archive.open("word/document.xml") as document_xml:
                    root = ET.parse(document_xml).getroot()
        except (KeyError, BadZipFile, ET.ParseError) as exc:
            raise ValueError(f"无法解析 docx 文档: {str(exc)}") from exc

        for paragraph in root.findall(".//w:p", namespace):
            text_nodes = [node.text or "" for node in paragraph.findall(".//w:t", namespace)]
            paragraph_text = "".join(text_nodes).strip()
            if paragraph_text:
                parts.append(paragraph_text)
        return "\n".join(parts)

    def _extract_pdf_text(self, content: bytes) -> str:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise ValueError("当前运行环境未安装 pypdf，暂时无法解析 PDF 文档") from exc

        reader = PdfReader(BytesIO(content))
        parts: List[str] = []
        for page in reader.pages:
            text = (page.extract_text() or "").strip()
            if text:
                parts.append(text)
        return "\n".join(parts)

    def _normalize_document_text(self, text: str) -> str:
        normalized_lines = []
        for raw_line in (text or "").splitlines():
            line = raw_line.strip()
            if line:
                normalized_lines.append(line)
        return "\n".join(normalized_lines).strip()
